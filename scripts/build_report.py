"""Dựng báo cáo đồ án và tài liệu hướng dẫn sử dụng (Word + PDF) từ Markdown, theo mẫu báo cáo thực hành của nhóm (T29).

    pip install -r scripts/requirements-docs.txt
    python scripts/build_report.py                 # cả hai tài liệu → report/*.docx
    python scripts/build_report.py baocao --pdf    # chỉ báo cáo, kèm PDF (in bằng Chrome/Edge headless)
    python scripts/build_report.py hdsd --pdf      # chỉ hướng dẫn sử dụng

Mẫu: report/template/mau-bao-cao-uit.docx (tách từ báo cáo Lab 04 – giữ nguyên bìa, checklist, trang nhận xét của giảng viên,
style, header / footer "Trang X/Y"). Nội dung sinh theo đúng quy ước của mẫu:
  • Tiêu đề cấp 1 "CHƯƠNG n: …" / "PHẦN n: …" sang trang mới; cấp 2–3 đánh số "1.1." / "1.1.1.".
  • Đoạn văn canh đều, thụt đầu dòng 0,5 cm, giãn dòng 1,15; gạch đầu dòng "•", danh sách số "1.".
  • Bảng: hàng tiêu đề nền xanh đậm chữ trắng, viền xám; khối mã: hộp nền xám, Consolas 10,5, tô màu cú pháp kiểu SSMS.
    Khối mã có dòng tiêu đề như mẫu ("Mã lệnh – …", "Kiểm thử – …") khi viết ```sql title="Mã lệnh – sp_DangKyHoc"```.
  • Hình: căn giữa, chú thích "Hình x.y. …" in nghiêng màu xám ngay dưới (lấy từ dòng "*Hình x.y – …*" trong Markdown).
  • Bảng: chú thích "Bảng x.y. …" ngay trên bảng (dòng "*Bảng x.y – …*" đứng trước bảng; phụ lục dùng "Bảng C.1").
    Script cảnh báo bảng / hình thiếu chú thích và số thứ tự nhảy cóc.
  • Mục lục, Danh mục hình ảnh, Danh mục bảng là trường TOC – Word hỏi cập nhật khi mở file, chọn Yes (hoặc Ctrl+A rồi
    F9); Danh mục từ viết tắt sinh từ bảng VIET_TAT (chỉ giữ từ có xuất hiện trong tài liệu).
  • Phần đầu (bìa → danh mục) là một section riêng: bìa không có tiêu đề chạy và số trang, các trang sau đánh số
    i, ii, iii…; thân bài đánh số lại từ 1, chân trang "Trang X/Y" với Y là số trang thân bài.
  • Giữ tiêu đề, câu dẫn ("…:"), chú thích bảng đi cùng nội dung phía sau; bảng ngắn (≤ 12 dòng) và khối mã ngắn
    (≤ 30 dòng) không bị cắt ngang trang; tên định danh trong bảng chỉ xuống dòng sau dấu "_".
  • <sub>…</sub>, <sup>…</sup> trong Markdown thành chỉ số dưới / trên.
Bản PDF dựng từ cùng nội dung qua HTML + CSS in ấn (mục lục là liên kết, không kèm số trang). Bản PDF nộp bài xuất bằng
Word: scripts/export_pdf.sh (macOS) hoặc scripts/export_pdf.ps1 (Windows).
Liên kết tới file trong repo được đổi thành đường dẫn GitHub.
"""
import argparse
import base64
import copy
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from posixpath import normpath

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Cm, Pt, RGBColor, Twips
from markdown_it import MarkdownIt
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "report"
TEMPLATE = REPORT / "template" / "mau-bao-cao-uit.docx"
GITHUB = "https://github.com/nhutruong-uit/ecms/blob/main/"
NGAY = "TP. Hồ Chí Minh, tháng 11 năm 2026"
TEN_HT = "HỆ THỐNG QUẢN LÝ TRUNG TÂM TIẾNG ANH (ECMS)"

DOCS = {
    "baocao": dict(
        src=[REPORT / f for f in ("01-tong-quan.md", "02-to-chuc-thong-tin.md", "03-xu-ly-thong-tin.md",
                                  "04-trinh-bay-thong-tin.md", "05-csdl-tien-tien.md", "06-ket-luan.md",
                                  "07-phu-luc.md") if (REPORT / f).exists()],
        out=REPORT / "BaoCao_IE103_Nhom1.docx",
        loai="BÁO CÁO ĐỒ ÁN", tieu_de=f"ĐỒ ÁN: {TEN_HT}", header=f"BÁO CÁO ĐỒ ÁN - {TEN_HT}",
        checklist=["Chương 1: Tổng quan – khảo sát nghiệp vụ, đặc tả yêu cầu",
                   "Chương 2: Tổ chức thông tin – ERD, lược đồ quan hệ, cài đặt CSDL",
                   "Chương 3: Xử lý thông tin – truy vấn, an ninh, lập trình CSDL",
                   "Chương 4: Trình bày thông tin – ứng dụng web, báo cáo",
                   "Chương 5: Mô hình CSDL tiên tiến – hướng đối tượng, phân tán, NoSQL",
                   "Kiểm thử tích hợp, CI và tài liệu hướng dẫn sử dụng"],
        nhan_xet=True, ma_consolas=True),
    "hdsd": dict(
        src=[ROOT / "docs" / "user-guide" / "huong-dan-su-dung.md"],
        out=REPORT / "HuongDanSuDung_ECMS.docx",
        loai="HƯỚNG DẪN SỬ DỤNG", tieu_de=TEN_HT, header=TEN_HT.removesuffix(" (ECMS)"),
        checklist=None, nhan_xet=False,
        ma_consolas=False),     # `mã` trong đoạn văn: cùng font Times New Roman, chỉ khối mã dùng Consolas
}

NAVY, BLUE, GREY = "1f3864", "2e75b6", "555555"
ZWSP = "\u200b"                     # chỗ được phép xuống dòng (vô hình) – chèn sau "_" của tên định danh trong bảng
BANG_NGAN, MA_NGAN = 12, 30         # bảng ≤ 12 dòng, khối mã ≤ 30 dòng: không cắt ngang trang
CHU_THICH_RE = re.compile(r"^\*(Hình|Bảng) ((?:\d+|[A-Z])\.\d+)")

# Danh mục từ viết tắt – chỉ in những từ thực sự xuất hiện trong tài liệu
VIET_TAT = {
    "2PC": "Two-Phase Commit – cam kết hai pha", "2PL": "Two-Phase Locking – khoá hai pha",
    "BASE": "Basically Available, Soft state, Eventually consistent", "BCNF": "Boyce–Codd Normal Form – dạng chuẩn Boyce–Codd",
    "CAP": "Consistency, Availability, Partition tolerance", "CI": "Continuous Integration – tích hợp liên tục",
    "CSDL": "Cơ sở dữ liệu", "CSV": "Comma-Separated Values", "DDL": "Data Definition Language – ngôn ngữ định nghĩa dữ liệu",
    "DDM": "Dynamic Data Masking – che dữ liệu động", "DIKW": "Data – Information – Knowledge – Wisdom",
    "DML": "Data Manipulation Language – ngôn ngữ thao tác dữ liệu", "DTD": "Document Type Definition",
    "ERD": "Entity Relationship Diagram – sơ đồ thực thể – mối kết hợp", "ETL": "Extract – Transform – Load",
    "FK": "Foreign Key – khoá ngoại", "FR": "Functional Requirement – yêu cầu chức năng",
    "HDSD": "Hướng dẫn sử dụng", "JSON": "JavaScript Object Notation", "KPI": "Key Performance Indicator – chỉ số hiệu suất",
    "MSDTC": "Microsoft Distributed Transaction Coordinator", "NFR": "Non-Functional Requirement – yêu cầu phi chức năng",
    "NoSQL": "Not only SQL – CSDL phi quan hệ", "ODBC": "Open Database Connectivity", "ODL": "Object Definition Language",
    "OID": "Object Identifier – định danh đối tượng", "OODB": "Object-Oriented Database – CSDL hướng đối tượng",
    "OQL": "Object Query Language", "ORDBMS": "Object-Relational DBMS – hệ QTCSDL quan hệ – đối tượng",
    "PDF": "Portable Document Format", "PITR": "Point-In-Time Recovery – phục hồi theo thời điểm",
    "PK": "Primary Key – khoá chính", "RBTV": "Ràng buộc toàn vẹn", "RDBMS": "Relational DBMS – hệ QTCSDL quan hệ",
    "RPO": "Recovery Point Objective – lượng dữ liệu tối đa được phép mất", "RTO": "Recovery Time Objective – thời gian khôi phục tối đa",
    "SQL": "Structured Query Language", "SSMS": "SQL Server Management Studio", "TDE": "Transparent Data Encryption – mã hoá dữ liệu trong suốt",
    "TLS": "Transport Layer Security – mã hoá đường truyền", "UML": "Unified Modeling Language", "XML": "eXtensible Markup Language",
    "XSD": "XML Schema Definition",
}


def viet_tat_md(text: str) -> str:
    dung = [k for k in sorted(VIET_TAT, key=str.lower) if re.search(rf"(?<![\w-]){re.escape(k)}(?![\w-])", text)]
    return "| Từ viết tắt | Ý nghĩa |\n|---|---|\n" + "".join(f"| {k} | {VIET_TAT[k]} |\n" for k in dung)
W_TRANG = 9070                     # twips – A4 trừ lề 2,5 cm hai bên (như mẫu)
CAP_1 = {"tài liệu tham khảo", "phụ lục"}


# ============================================================================= tô màu cú pháp (dùng chung cho docx và HTML)
SQL_KW = set("""ADD ALL ALTER AND ANY AS ASC AUTHORIZATION BACKUP BEGIN BETWEEN BREAK BY CASCADE CASE CATCH CHECK CLOSE CLUSTERED
COLLATE COLUMN COMMIT CONSTRAINT CONTINUE CREATE CROSS CURSOR DATABASE DEALLOCATE DECLARE DEFAULT DELETE DENY DESC DISTINCT DROP
ELSE END EXEC EXECUTE EXISTS FETCH FILE FOR FOREIGN FROM FULL FUNCTION GO GOTO GRANT GROUP HAVING IDENTITY IF IN INDEX INNER INSERT
INTERSECT INTO IS JOIN KEY LEFT LIKE LOGIN MERGE NEXT NOCOUNT NOT NULL OF OFF ON OPEN OPTION OR ORDER OUTER OUTPUT OVER PARTITION
PERCENT PRIMARY PROC PROCEDURE RAISERROR REFERENCES RESTORE RETURN RETURNS REVOKE RIGHT ROLE ROLLBACK ROWS SAVE SCHEMA SCHEMABINDING
SELECT SET STATUS TABLE THEN THROW TIES TO TOP TRAN TRANSACTION TRIGGER TRY TRUNCATE UNION UNIQUE UPDATE USE USER VALUES VIEW
WHEN WHERE WHILE WITH INSTEAD AFTER NOCHECK READONLY OFFSET ONLY APPLY PIVOT EXCEPT ROLLUP CUBE MAX INT BIGINT NVARCHAR VARCHAR CHAR
NCHAR DATE DATETIME DATETIME2 DECIMAL NUMERIC BIT FLOAT XML TINYINT SMALLINT MONEY VARBINARY PERSISTED NONCLUSTERED INCLUDE""".split())
SQL_FN = set("""COUNT SUM AVG MIN MAX ISNULL COALESCE CAST CONVERT GETDATE SYSDATETIME DATEADD DATEDIFF DATEPART YEAR MONTH DAY
STRING_AGG CONCAT LEN LEFT RIGHT SUBSTRING UPPER LOWER ROUND ROW_NUMBER RANK DENSE_RANK NTILE LAG LEAD DB_ID OBJECT_ID OBJECT_NAME
SESSION_CONTEXT SUSER_SNAME HASHBYTES NEWID IIF CHOOSE FORMAT OPENJSON JSON_VALUE ERROR_MESSAGE ERROR_NUMBER SCOPE_IDENTITY
TRIM REPLACE ABS CEILING FLOOR EOMONTH DATENAME DATEFROMPARTS""".split())
PY_KW = set("def class return import from as if elif else for while in not and or is None True False with try except raise lambda pass".split())
SH_KW = set("cd export source docker sudo pip python python3 git streamlit pytest cp mv rm mkdir if then fi for do done echo set".split())

TOKEN_RE = {
    "sql": re.compile(r"(--[^\n]*)|(N?'(?:[^']|'')*')|(@@?\w+)|([A-Za-z_#][\w#$]*)|(\d+(?:\.\d+)?)|(\s+|.)"),
    "py": re.compile(r"(#[^\n]*)|(\"[^\"\n]*\"|'[^'\n]*')|(@\w+)|([A-Za-z_]\w*)|(\d+(?:\.\d+)?)|(\s+|.)"),
}


def to_mau(code: str, lang: str) -> list[list[tuple[str, str]]]:
    """Mỗi dòng → [(màu hex, đoạn chữ)] theo bảng màu SSMS / VS Code."""
    lang = (lang or "").lower()
    kind = "sql" if lang in ("sql", "tsql", "t-sql") else "py" if lang in ("python", "py", "bash", "sh", "shell", "powershell",
                                                                             "ps1", "zsh", "yaml", "yml", "dotenv", "env") else None
    kws = PY_KW if lang in ("python", "py") else SH_KW
    lines = []
    for ln in code.rstrip("\n").split("\n"):
        if kind is None:
            lines.append([("000000", ln)])
            continue
        out = []
        for m in TOKEN_RE[kind].finditer(ln):
            com, s, var, word, num, other = m.groups()
            if com:
                out.append(("008000", com))
            elif s:
                out.append(("a31515", s))
            elif var:
                out.append(("000000", var))
            elif word:
                if kind == "sql":
                    u = word.upper()
                    out.append(("ff00ff" if u in SQL_FN else "0000ff" if u in SQL_KW else "000000", word))
                else:
                    out.append(("0000ff" if word in kws else "000000", word))
            elif num:
                out.append(("000000", num))
            else:
                out.append(("000000", other))
        lines.append(out or [("000000", "")])
    return lines


# ============================================================================= chuẩn hoá tiêu đề, chú thích
def tieu_de(level: int, text: str) -> tuple[int, str]:
    text = text.replace("**", "").replace("`", "").strip()
    if text.lower() in CAP_1:
        return 1, text.upper()
    if level == 1:
        m = re.match(r"^(Chương|Phần)\s+(\d+)\s*[.:–-]\s*(.+)$", text, re.I)
        return 1, (f"{m[1].upper()} {m[2]}: {m[3].upper()}" if m else text.upper())
    m = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.+)$", text)
    return level, (f"{m[1]}. {m[2]}" if m else text)


def chu_thich(raw: str) -> str:
    """'*Hình 4.1 – Mô tả*' → 'Hình 4.1. Mô tả' (giữ định dạng inline còn lại)."""
    t = raw.strip()
    if t.startswith("*") and t.endswith("*"):
        t = t[1:-1]
    return re.sub(r"^(Hình|Bảng) ((?:\d+|[A-Z])\.\d+)\s*[–:.-]\s*", r"\1 \2. ", t)


def fence_info(info: str) -> "tuple[str, str | None]":
    """'sql title="Mã lệnh – sp_DangKyHoc"' → ('sql', 'Mã lệnh – sp_DangKyHoc'): dòng tiêu đề của khối mã như mẫu Lab 04."""
    info = (info or "").strip()
    m = re.search(r'title="([^"]+)"', info)
    return ("" if info.startswith("title=") else info.split(" ")[0]), (m[1] if m else None)


def link_url(src: Path, href: str) -> str:
    if re.match(r"^[a-z]+:", href) or href.startswith("#"):
        return href
    return GITHUB + normpath((src.parent.relative_to(ROOT) / href).as_posix())


def md_parser():
    return MarkdownIt("commonmark", {"html": True}).enable("table")


# ============================================================================= OOXML helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def run_fmt(r, font="Times New Roman", size=None, bold=None, italic=None, color=None, underline=False):
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = el("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), font)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color.upper())
    if underline:
        r.underline = True
    return r


def p_fmt(p, align="both", after=120, before=0, line=276, first=None, left=None, hanging=None, keep_next=False):
    pf = p.paragraph_format
    pf.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.space_after, pf.space_before = Twips(after), Twips(before)
    pf.line_spacing = line / 240
    if first is not None:
        pf.first_line_indent = Twips(first)
    if left is not None:
        pf.left_indent = Twips(left)
    if hanging is not None:
        pf.first_line_indent = Twips(-hanging)
    if keep_next:
        pf.keep_with_next = True
    return p


def numpr(p, num_id: int, ilvl: int = 0):
    ppr = p._p.get_or_add_pPr()
    n = el("w:numPr")
    n.append(el("w:ilvl", w_val=ilvl))
    n.append(el("w:numId", w_val=num_id))
    ppr.insert(0 if ppr.find(qn("w:pStyle")) is None else 1, n)


def field(p, instr: str, placeholder: str):
    def fc(t):
        r = el("w:r")
        r.append(el("w:fldChar", w_fldCharType=t))
        return r
    p._p.append(fc("begin"))
    r = el("w:r")
    it = el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r.append(it)
    p._p.append(r)
    p._p.append(fc("separate"))
    run_fmt(p.add_run(placeholder), italic=True, color=GREY)
    p._p.append(fc("end"))


def hyperlink(p, url, text, bold=False, italic=False, code=False, size=None):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = el("w:hyperlink", r_id=rid)
    r = el("w:r")
    rpr = el("w:rPr")
    font = "Consolas" if code else "Times New Roman"
    rpr.append(el("w:rFonts", w_ascii=font, w_hAnsi=font, w_eastAsia=font, w_cs=font))
    if bold:
        rpr.append(el("w:b"))
    if italic:
        rpr.append(el("w:i"))
    rpr.append(el("w:color", w_val=BLUE))
    if size:
        rpr.append(el("w:sz", w_val=int(size * 2)))
    r.append(rpr)
    t = el("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def cell_fmt(tc, fill=None, border="bbbbbb", mar=(60, 100), width=None, valign="center"):
    pr = tc._tc.get_or_add_tcPr()
    if width:
        pr.append(el("w:tcW", w_w=width, w_type="dxa"))
    b = el("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b.append(el(f"w:{side}", w_val="single", w_sz=4, w_space=0, w_color=border))
    pr.append(b)
    if fill:
        pr.append(el("w:shd", w_val="clear", w_color="auto", w_fill=fill))
    m = el("w:tcMar")
    for side, v in (("top", mar[0]), ("left", mar[1]), ("bottom", mar[0]), ("right", mar[1])):
        m.append(el(f"w:{side}", w_w=v, w_type="dxa"))
    pr.append(m)
    if valign:
        pr.append(el("w:vAlign", w_val=valign))


def table_fmt(tbl, widths):
    pr = tbl._tbl.tblPr
    for tag in ("w:tblStyle", "w:tblW", "w:tblLook"):
        for e in pr.findall(qn(tag)):
            pr.remove(e)
    pr.append(el("w:tblW", w_w=sum(widths), w_type="dxa"))
    pr.append(el("w:jc", w_val="left"))
    b = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(el(f"w:{side}", w_val="single", w_sz=4, w_space=0, w_color="000000"))
    pr.append(b)
    pr.append(el("w:tblLayout", w_type="fixed"))
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))


def cot_rong(texts, weights, size) -> list[int]:
    """Độ rộng cột (twips): mỗi cột đủ chỗ cho từ dài nhất của nó (không ngắt giữa từ như "5000 / 1"),
    phần còn lại chia theo độ dài nội dung."""
    def rong(w):                          # twips; Times New Roman: chữ hoa, số ≈ 0,72 em, chữ thường ≈ 0,5 em
        return sum(0.72 if c.isupper() or c.isdigit() or c in "()_%" else 0.5 for c in w) * size * 20
    # lề ô 2 × 100 twips + dư; hàng đầu (tiêu đề in đậm) rộng hơn khoảng 10%
    toi_thieu = [int(min(max([rong(w) * (1.1 if k == 0 else 1) for k, t in enumerate(col) for w in t.split()]
                             or [size * 10]) + 380, W_TRANG * 3 // 10))
                 for col in texts]
    if sum(toi_thieu) >= W_TRANG:         # quá nhiều cột: chia theo tỷ lệ độ rộng tối thiểu
        return [int(W_TRANG * m / sum(toi_thieu)) for m in toi_thieu]
    co_dinh: set[int] = set()             # cột bị giữ ở mức tối thiểu; phần còn lại chia theo tỷ lệ cho các cột khác
    while True:
        con_lai = W_TRANG - sum(toi_thieu[i] for i in co_dinh)
        tong = sum(w for i, w in enumerate(weights) if i not in co_dinh)
        widths = [toi_thieu[i] if i in co_dinh else int(con_lai * w / tong) for i, w in enumerate(weights)]
        thieu = {i for i, (w, m) in enumerate(zip(widths, toi_thieu)) if w < m} - co_dinh
        if not thieu:
            return widths
        co_dinh |= thieu


def spacer(doc):
    p_fmt(doc.add_paragraph(), align="left", after=60, line=240)


# ============================================================================= Markdown → docx
class DocxWriter:
    def __init__(self, doc, src: Path, ma_consolas=True):
        self.doc, self.src, self.ma_consolas = doc, src, ma_consolas
        self.num_id = 100

    # ----- inline
    def inline(self, p, children, bold=False, italic=False, size=None, color=None, ngat=False):
        """ngat=True (ô bảng): cho phép xuống dòng sau "_" để tên định danh dài không bị Word cắt giữa chữ."""
        b0, i0, href, vt = bold, italic, None, None
        for t in children or []:
            if t.type == "strong_open":
                bold = True
            elif t.type == "strong_close":
                bold = b0
            elif t.type == "em_open":
                italic = True
            elif t.type == "em_close":
                italic = i0
            elif t.type == "link_open":
                href = link_url(self.src, t.attrGet("href") or "")
            elif t.type == "link_close":
                href = None
            elif t.type == "html_inline":
                tag = t.content.strip().lower()
                if tag in ("<sub>", "<sup>"):
                    vt = tag[1:4]
                elif tag in ("</sub>", "</sup>"):
                    vt = None
                elif tag in ("<br>", "<br/>", "<br />"):
                    p.add_run().add_break()
            elif t.type in ("text", "code_inline"):
                code = t.type == "code_inline"
                mono = code and self.ma_consolas
                csize = (size or 13) - 1 if mono else size
                text = t.content.replace("_", "_" + ZWSP) if ngat else t.content
                if href and not href.startswith("#"):
                    hyperlink(p, href, text, bold, italic, mono, csize)
                else:
                    r = run_fmt(p.add_run(text), "Consolas" if mono else "Times New Roman", csize, bold, italic,
                                color or (NAVY if code else None))
                    if vt:
                        r.font.subscript, r.font.superscript = vt == "sub", vt == "sup"
            elif t.type == "softbreak":
                p.add_run(" ")
            elif t.type == "hardbreak":
                p.add_run().add_break()

    # ----- blocks
    def heading(self, level, text):
        level, text = tieu_de(level, text)
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        if level == 1:
            p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True          # tiêu đề không nằm một mình cuối trang
        run_fmt(p.add_run(text), size=13)

    def paragraph(self, children, raw):
        p = self.doc.add_paragraph()
        dan = raw.rstrip().endswith(":")                  # câu dẫn của bảng / danh sách / khối mã: đi cùng phần sau
        p_fmt(p, first=284, keep_next=dan)                 # mọi đoạn thân bài thụt 0,5 cm, kể cả câu dẫn
        self.inline(p, children)

    def image(self, src):
        path = (self.src.parent / src).resolve()
        with Image.open(path) as im:
            w, h = im.size
        width = 15.0 if w >= 700 else min(15.0, w / 96 * 2.54)
        if width * h / w > 20.0:
            width = 20.0 * w / h
        p = self.doc.add_paragraph()
        p_fmt(p, align="center", after=0, line=240, keep_next=True)
        p.add_run().add_picture(str(path), width=Cm(width))

    def caption(self, raw, figure=True):
        p = self.doc.add_paragraph(style="FigureCaption" if figure else "TableCaption")
        p_fmt(p, align="center", after=160 if figure else 60, before=160 if figure else 120, line=240, first=0,
              keep_next=not figure)
        self.inline(p, md_parser().parseInline(chu_thich(raw))[0].children, italic=True, color=GREY)

    def table(self, rows, aligns):
        ncol = max(len(r) for r in rows)
        size = 12 if ncol <= 4 else 11 if ncol <= 6 else 10
        texts = [["".join(c.content for c in (r[i] or []) if c.type in ("text", "code_inline")) if i < len(r) else ""
                  for r in rows] for i in range(ncol)]
        # trọng số theo độ dài trung bình (cột nhiều chữ được rộng hơn) cộng một phần độ dài lớn nhất
        weights = [min(sum(len(t) for t in col[1:]) / max(len(col) - 1, 1), 150) + 0.3 * min(max(len(t) for t in col), 60) + 3
                   for col in texts]
        widths = cot_rong(texts, weights, size)
        tbl = self.doc.add_table(rows=len(rows), cols=ncol)
        table_fmt(tbl, widths)
        ngan = len(rows) <= BANG_NGAN
        for ri, cells in enumerate(rows):
            tr = tbl.rows[ri]._tr
            trpr = tr.get_or_add_trPr()
            trpr.append(el("w:cantSplit", w_val=0))
            trpr.append(el("w:tblHeader", w_val=1 if ri == 0 else 0))
            for ci in range(ncol):
                cell = tbl.cell(ri, ci)
                cell_fmt(cell, fill=NAVY if ri == 0 else None, width=widths[ci])
                p = cell.paragraphs[0]
                a = "center" if ri == 0 else (aligns[ci] if ci < len(aligns) else "left")
                # keep_next trên mọi dòng trừ dòng cuối: Word giữ cả bảng ngắn trên một trang; bảng dài vẫn giữ
                # tiêu đề cột đi cùng ít nhất một dòng dữ liệu
                p_fmt(p, align=a, after=0, line=252, keep_next=(ngan and ri < len(rows) - 1) or ri == 0)
                if ci < len(cells):
                    self.inline(p, cells[ci], bold=(ri == 0), size=size, color="ffffff" if ri == 0 else None, ngat=True)
        spacer(self.doc)

    def code(self, text, info):
        lang, title = fence_info(info)
        tbl = self.doc.add_table(rows=2 if title else 1, cols=1)
        table_fmt(tbl, [W_TRANG])
        if title:
            head = tbl.cell(0, 0)
            cell_fmt(head, fill="e7ecf5", border="bfbfbf", mar=(40, 140), width=W_TRANG)
            p = head.paragraphs[0]
            p_fmt(p, align="left", after=0, line=240, keep_next=True)
            run_fmt(p.add_run(title), size=11, bold=True, color=NAVY)
            tbl.rows[0]._tr.get_or_add_trPr().append(el("w:cantSplit", w_val=1))
        cell = tbl.cell(1 if title else 0, 0)
        cell_fmt(cell, fill="f7f7f7", border="bfbfbf", mar=(100, 140), width=W_TRANG, valign=None)
        lines = to_mau(text, lang)
        if len(lines) <= MA_NGAN:                         # khối mã ngắn: không cắt ngang trang
            tbl.rows[-1]._tr.get_or_add_trPr().append(el("w:cantSplit", w_val=1))
        for k, parts in enumerate(lines):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            p_fmt(p, align="left", after=0, line=240)
            for color, s in parts:
                run_fmt(p.add_run(s), "Consolas", 10.5, color=color)
        spacer(self.doc)

    def new_list(self, ordered):
        if not ordered:
            return 2
        numbering = self.doc.part.numbering_part.element
        # numId lớn nhất đang có + 1: mỗi chương dùng một DocxWriter riêng, đếm từ 100 ở mỗi chương thì danh sách
        # của các chương trùng numId và Word đánh số nối tiếp (17, 18… thay vì 1, 2…)
        self.num_id = max([int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))] + [100]) + 1
        n = el("w:num", w_numId=self.num_id)
        n.append(el("w:abstractNumId", w_val=1))
        o = el("w:lvlOverride", w_ilvl=0)
        o.append(el("w:startOverride", w_val=1))
        n.append(o)
        numbering.append(n)
        return self.num_id

    def list_block(self, tokens, i, depth=0):
        ordered = tokens[i].type == "ordered_list_open"
        close = tokens[i].type.replace("_open", "_close")
        nid = self.new_list(ordered)
        left = 567 + 567 * depth
        j, first = i + 1, False
        while tokens[j].type != close:
            t = tokens[j]
            if t.type == "list_item_open":
                first = True
            elif t.type == "paragraph_open":
                p = self.doc.add_paragraph()
                if first:
                    p_fmt(p, after=80, left=left, hanging=283 if ordered else 300)
                    numpr(p, nid, min(depth, 1) if not ordered else 0)
                    first = False
                else:
                    p_fmt(p, after=80, left=left, first=0)
                self.inline(p, tokens[j + 1].children)
                j += 2
            elif t.type in ("bullet_list_open", "ordered_list_open"):
                j = self.list_block(tokens, j, depth + 1)
            elif t.type == "fence":
                self.code(t.content, t.info)
            j += 1
        return j

    def canh_bao(self, line, msg):
        print(f"  ⚠ {self.src.name}:{line + 1}: {msg}", file=sys.stderr)

    def kiem_so(self, raw, line):
        """Số chú thích phải tăng dần 1 trong mỗi chương / phụ lục."""
        m = CHU_THICH_RE.match(raw)
        if not m:
            return
        loai, (chuong, so) = m[1], m[2].split(".")
        truoc = self.so.get((loai, chuong), 0)
        if int(so) != truoc + 1:
            self.canh_bao(line, f"{loai} {m[2]} – số trước đó là {chuong}.{truoc}")
        self.so[(loai, chuong)] = int(so)

    def write(self, text):
        tokens = md_parser().parse(text)
        i, pending, co_chu_thich = 0, None, False
        self.so = {}
        while i < len(tokens):
            t = tokens[i]
            if t.type == "heading_open":
                self.heading(int(t.tag[1]), tokens[i + 1].content)
                i += 3
                continue
            if t.type == "paragraph_open":
                inl = tokens[i + 1]
                raw, kids = inl.content.strip(), inl.children or []
                line = (t.map or [0])[0]
                if re.match(r"^\*Bảng [\dA-Z]", raw) and i + 3 < len(tokens) and tokens[i + 3].type == "table_open":
                    self.kiem_so(raw, line)
                    self.caption(raw, figure=False)
                    co_chu_thich = True
                elif re.match(r"^\*Hình [\dA-Z]", raw):
                    self.kiem_so(raw, line)
                    self.caption(raw)
                elif len(kids) == 1 and kids[0].type == "image":
                    self.image(kids[0].attrGet("src"))
                    nxt = tokens[i + 4].content.strip() if i + 4 < len(tokens) and tokens[i + 3].type == "paragraph_open" else ""
                    if not re.match(r"^\*Hình ", nxt):
                        self.canh_bao(line, "hình chưa có chú thích *Hình x.y – …*")
                else:
                    self.paragraph(kids, raw)
                i += 3
                continue
            if t.type == "table_open":
                if not co_chu_thich and getattr(self, "bat_buoc_chu_thich", True):
                    self.canh_bao((t.map or [0])[0], "bảng chưa có chú thích *Bảng x.y – …*")
                co_chu_thich = False
                rows, aligns, j = [], [], i + 1
                while tokens[j].type != "table_close":
                    tj = tokens[j]
                    if tj.type == "tr_open":
                        rows.append([])
                    elif tj.type in ("th_open", "td_open") and len(rows) == 1:
                        st = tj.attrGet("style") or ""
                        aligns.append("right" if "right" in st else "center" if "center" in st else "left")
                    elif tj.type == "inline":
                        rows[-1].append(tj.children)
                    j += 1
                self.table(rows, aligns)
                i = j + 1
                continue
            if t.type in ("bullet_list_open", "ordered_list_open"):
                i = self.list_block(tokens, i) + 1
                continue
            if t.type in ("fence", "code_block"):
                self.code(t.content, t.info or "")
            i += 1


def set_text(p_el, text):
    first = True
    for r in p_el.findall(qn("w:r")):
        ts = r.findall(qn("w:t"))
        if not ts:
            continue
        if first:
            ts[0].text = text
            for t in ts[1:]:
                r.remove(t)
            first = False
        else:
            for t in ts:
                r.remove(t)


def chen_truong(p, instr, text, color="888888"):
    """Trường đơn giản (PAGE, SECTIONPAGES…) với giá trị tạm – Word tính lại khi mở / in."""
    for kind in ("begin", None, "separate", "end"):
        r = el("w:r")
        rpr = el("w:rPr")
        rpr.append(el("w:color", w_val=color))
        r.append(rpr)
        if kind is None:
            it = el("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            r.append(it)
        else:
            r.append(el("w:fldChar", w_fldCharType=kind))
            if kind == "separate":
                t = el("w:t")
                t.text = text
                r2 = el("w:r")
                r2.append(copy.deepcopy(rpr))
                r2.append(t)
                p._p.append(r)
                p._p.append(r2)
                continue
        p._p.append(r)


def tach_phan_dau(doc):
    """Tách bìa → danh mục thành section riêng: bìa không tiêu đề chạy / số trang, các trang sau đánh số i, ii…;
    thân bài (từ "CHƯƠNG 1" / "PHẦN 1") đánh số lại từ 1, chân trang "Trang X/Y" với Y = số trang của thân bài."""
    body = doc.element.body
    dau = next((p for p in body.iter(qn("w:p"))
                if (s := p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")) is not None and s.get(qn("w:val")) in ("Heading1", "Heading 1")
                and re.match(r"^(CHƯƠNG|PHẦN) 1\b", "".join(t.text or "" for t in p.iter(qn("w:t"))))), None)
    if dau is None:
        return
    cuoi = body.sectPr
    sp = copy.deepcopy(cuoi)
    for ref in sp.findall(qn("w:footerReference")):
        sp.remove(ref)
    sp.find(qn("w:pgNumType")).set(qn("w:fmt"), "lowerRoman")
    ngat = el("w:p")
    ppr = el("w:pPr")
    ppr.append(sp)
    ngat.append(ppr)
    dau.addprevious(ngat)
    pf = dau.find(qn("w:pPr")).find(qn("w:pageBreakBefore"))
    if pf is not None:                                    # section mới đã sang trang – bỏ ngắt trang thừa
        dau.find(qn("w:pPr")).remove(pf)
    phan_dau = Section(sp, doc.part)
    phan_dau.different_first_page_header_footer = True     # bìa: không tiêu đề chạy, không số trang
    phan_dau.first_page_header.is_linked_to_previous = False
    phan_dau.first_page_footer.is_linked_to_previous = False
    phan_dau.footer.is_linked_to_previous = False
    p = phan_dau.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chen_truong(p, "PAGE", "ii")
    for it in doc.sections[-1].footer._element.iter(qn("w:instrText")):
        if it.text.strip() == "NUMPAGES":
            it.text = "SECTIONPAGES"


def build_docx(cfg) -> Path:
    doc = Document(TEMPLATE)
    body = doc.element.body
    paras = {("".join(t.text or "" for t in p.iter(qn("w:t")))): p for p in body.iter(qn("w:p"))}
    # bìa: bớt 3 dòng trống để dòng ngày không bị đẩy sang trang 2
    trong = [e for e in list(body)[:body.index(paras["«NGAY»"])] if e.tag == qn("w:p") and not "".join(e.itertext()).strip()]
    for e in trong[-3:]:
        body.remove(e)
    doc.styles["FigureCaption"].font.size = Pt(13)          # mẫu để 11 pt; chú thích hình / bảng cùng cỡ thân bài
    # Ký hiệu "•" của danh sách không khai báo font → Word lấy Arial; mục lục Word sinh ra gắn font theme (Calibri)
    # vào dấu chấm dẫn → đặt rõ Times New Roman cho cả hai
    for lvl in doc.part.numbering_part.element.iter(qn("w:lvl")):
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = el("w:rPr")
            lvl.append(rpr)
        for cu in rpr.findall(qn("w:rFonts")):
            rpr.remove(cu)
        rpr.insert(0, el("w:rFonts", w_ascii="Times New Roman", w_hAnsi="Times New Roman", w_cs="Times New Roman"))
    theme = next(r.target_part for r in doc.part.rels.values() if r.reltype.endswith("/theme"))
    theme._blob = re.sub(rb'(<a:(?:major|minor)Font><a:latin typeface=")[^"]*', rb"\1Times New Roman", theme.blob)
    if "TableCaption" not in [s.name for s in doc.styles]:
        tc = doc.styles.add_style("TableCaption", WD_STYLE_TYPE.PARAGRAPH)
        tc.base_style = doc.styles["FigureCaption"]
    set_text(paras["«LOAI»"], cfg["loai"])
    set_text(paras["«TIEU_DE»"], cfg["tieu_de"])
    set_text(paras["«NGAY»"], NGAY)
    for p in doc.sections[0].header.paragraphs:
        for r in p.runs:
            if "«HEADER»" in r.text:
                r.text = cfg["header"]
    tbls = [e for e in body if e.tag == qn("w:tbl")]
    ck_title, ck = paras["«CHECKLIST»"], tbls[1]
    if cfg["checklist"]:
        set_text(ck_title, "CHECKLIST TIẾN ĐỘ HOÀN THÀNH")
        hdr = ck.findall(qn("w:tr"))[0].findall(qn("w:tc"))[0].find(qn("w:p"))
        set_text(hdr, "Nội dung")                         # mẫu Lab ghi "Nội dung bài thực hành"
        proto = ck.findall(qn("w:tr"))[1]
        for noi_dung in cfg["checklist"]:
            row = copy.deepcopy(proto)
            tcs = row.findall(qn("w:tc"))
            set_text(tcs[0].find(qn("w:p")), noi_dung)
            set_text(tcs[-1].find(qn("w:p")), "X")
            ck.append(row)
        ck.remove(proto)
    else:
        body.remove(ck_title)
        body.remove(ck)
    if not cfg["nhan_xet"]:
        nx = paras["NHẬN XÉT CỦA GIẢNG VIÊN"]
        body.remove(nx.getnext())
        body.remove(nx)
    goi_y = "Mở bằng Microsoft Word và chọn Yes khi được hỏi cập nhật trường để hiện mục lục."
    for title, instr in (("MỤC LỤC", 'TOC \\h \\u \\z \\t "Heading 1,1,Heading 2,2,Heading 3,3,"'),
                         ("DANH MỤC HÌNH ẢNH", 'TOC \\h \\z \\t "FigureCaption,1"')):
        anchor = paras[title]
        tmp = doc.add_paragraph()
        field(tmp, instr, goi_y)
        anchor.addnext(tmp._p)
    noi_dung = "\n".join(src.read_text(encoding="utf-8") for src in cfg["src"])
    for title in ("DANH MỤC BẢNG", "DANH MỤC TỪ VIẾT TẮT"):  # cùng định dạng tiêu đề "DANH MỤC HÌNH ẢNH" của mẫu
        tmp = doc.add_paragraph()
        tp = copy.deepcopy(paras["DANH MỤC HÌNH ẢNH"])
        set_text(tp, title)
        tmp._p.addprevious(tp)
        if title == "DANH MỤC BẢNG":
            field(tmp, 'TOC \\h \\z \\t "TableCaption,1"', goi_y)
        else:
            body.remove(tmp._p)
            w = DocxWriter(doc, cfg["src"][0], cfg["ma_consolas"])
            w.bat_buoc_chu_thich = False
            w.write(viet_tat_md(noi_dung))
    for src in cfg["src"]:
        DocxWriter(doc, src, cfg["ma_consolas"]).write(src.read_text(encoding="utf-8"))
    tach_phan_dau(doc)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(el("w:updateFields", w_val="true"))
    doc.core_properties.title = cfg["tieu_de"]
    doc.core_properties.author = "Nhóm 1 – IE103.Q21.VB2"
    cfg["out"].parent.mkdir(parents=True, exist_ok=True)
    doc.save(cfg["out"])
    print(f"✓ {cfg['out'].relative_to(ROOT)}")
    return cfg["out"]


# ============================================================================= Markdown → HTML → PDF
CSS = """
@page { size: A4; margin: 2.54cm 2.5cm;
  @top-center { content: "%HEADER%"; font: 10.5pt 'Times New Roman'; color: #888; }
  @bottom-center { content: "Trang " counter(page) "/" counter(pages); font: 10.5pt 'Times New Roman'; color: #888; } }
/* phần đầu (bìa → danh mục): số trang La Mã; bìa không có tiêu đề chạy và số trang */
@page front { @bottom-center { content: "Trang " counter(page, lower-roman); font: 10.5pt 'Times New Roman'; color: #888; } }
@page front:first { @top-center { content: none; } @bottom-center { content: none; } }
.front { page: front; }
body { font: 13pt/1.15 'Times New Roman', serif; color: #000; margin: 0; }
.pb { break-before: page; }
h1 { font-size: 13pt; color: #1f3864; break-before: page; margin: 12pt 0 8pt; }
h2 { font-size: 13pt; color: #2e75b6; margin: 10pt 0 6pt; } h3, h4 { font-size: 13pt; margin: 8pt 0 5pt; }
h1, h2, h3, h4 { break-after: avoid; }
p { text-align: justify; margin: 0 0 6pt; text-indent: 0.5cm; } p.lead { break-after: avoid; }
li { text-align: justify; margin-bottom: 4pt; } li p { text-indent: 0; margin: 0; }
a { color: #2e75b6; text-decoration: none; }
code { font: 12pt Consolas, Menlo, monospace; color: #1f3864; }
table { border-collapse: collapse; width: 100%; margin: 0 0 3pt; font-size: 11.5pt; line-height: 1.05; }
th, td { border: 0.5pt solid #bbb; padding: 3pt 5pt; vertical-align: middle; overflow-wrap: break-word; }
/* Chrome thu nhỏ CẢ tài liệu khi một phần tử rộng hơn trang (tên dài như sp_TinhLuongGiaoVien_SetBased trong bảng
   nhiều cột) → chữ còn ~8pt. Trong ô bảng, tên định danh có ký tự ngắt vô hình sau mỗi "_" (xem ZWSP) nên xuống dòng
   đúng chỗ thay vì cắt giữa chữ ("NHANVIE N"); khối mã vẫn được ngắt ở bất kỳ đâu. */
td code, th code { overflow-wrap: break-word; } pre.code { overflow-wrap: anywhere; word-break: break-word; }
table.short, pre.code.short { break-inside: avoid; }
th { background: #1f3864; color: #fff; text-align: center; }
tr { break-inside: avoid; } thead { display: table-header-group; }
pre.code { background: #f7f7f7; border: 0.5pt solid #bfbfbf; padding: 5pt 7pt; font: 10.5pt/1.2 Consolas, Menlo, monospace;
  white-space: pre-wrap; margin: 0 0 6pt; }
.ctitle { background: #e7ecf5; border: 0.5pt solid #bfbfbf; border-bottom: 0; padding: 2pt 7pt; font: bold 11pt 'Times New Roman';
  color: #1f3864; break-after: avoid; } pre.code.titled { margin-top: 0; }
p.fig { text-align: center; text-indent: 0; margin: 6pt 0 0; break-after: avoid; } p.fig img { max-width: 15cm; max-height: 20cm; }
p.fcap, p.tcap { text-align: center; text-indent: 0; font-style: italic; color: #555; } p.fcap { margin: 8pt 0 8pt; }
p.tcap { margin: 6pt 0 3pt; break-after: avoid; }
.cover p { text-align: center; text-indent: 0; margin: 0 0 3pt; } .vt td:first-child { white-space: nowrap; font-weight: bold; } .cover table { margin-top: 6pt; font-size: 13pt; }
.title { font-weight: bold; color: #1f3864; text-align: center; margin: 0 0 9pt; }
.toc ul { list-style: none; padding: 0; margin: 0; } .toc li { margin: 0 0 5pt; text-align: left; }
.toc li.l2 { padding-left: 13pt; } .toc li.l3 { padding-left: 26pt; }
.nx td { height: 0.95cm; border-color: #bbb; }
"""


def build_pdf(cfg):
    md = md_parser()
    phan, muc, hinh, bang, n = [], [], [], [], 0
    for src in cfg["src"]:
        tokens = md.parse(src.read_text(encoding="utf-8"))
        out, i, trong_bang = [], 0, False
        while i < len(tokens):
            t = tokens[i]
            if t.type == "table_open":
                trong_bang = True
                rows = 0
                for k in range(i, len(tokens)):
                    rows += tokens[k].type == "tr_open"
                    if tokens[k].type == "table_close":
                        break
                if rows <= BANG_NGAN:
                    t.attrSet("class", "short")
            elif t.type == "table_close":
                trong_bang = False
            if t.type == "inline" and trong_bang:
                for c in t.children or []:
                    if c.type in ("text", "code_inline"):
                        c.content = c.content.replace("_", "_" + ZWSP)
            if t.type == "inline":
                for c in t.children or []:
                    if c.type == "link_open":
                        c.attrSet("href", link_url(src, c.attrGet("href") or ""))
                    elif c.type == "image":
                        c.attrSet("src", (src.parent / c.attrGet("src")).resolve().as_uri())
            if t.type == "heading_open":
                level, text = tieu_de(int(t.tag[1]), tokens[i + 1].content)
                n += 1
                out.append(_html(f'<h{min(level, 4)} id="h{n}">{html.escape(text)}</h{min(level, 4)}>\n'))
                if level <= 3:
                    muc.append((level, f"h{n}", text))
                i += 3
                continue
            if t.type == "paragraph_open":
                inl = tokens[i + 1]
                raw, kids = inl.content.strip(), inl.children or []
                if re.match(r"^\*Bảng [\dA-Z]", raw) and i + 3 < len(tokens) and tokens[i + 3].type == "table_open":
                    n += 1
                    bang.append((f"h{n}", html.unescape(re.sub(r"<[^>]+>", "", md.renderInline(chu_thich(raw))))))
                    out.append(_html(f'<p class="tcap" id="h{n}">{md.renderInline(chu_thich(raw))}</p>\n'))
                    i += 3
                    continue
                if re.match(r"^\*Hình [\dA-Z]", raw):
                    n += 1
                    hinh.append((f"h{n}", html.unescape(re.sub(r"<[^>]+>", "", md.renderInline(chu_thich(raw))))))
                    out.append(_html(f'<p class="fcap" id="h{n}">{md.renderInline(chu_thich(raw))}</p>\n'))
                    i += 3
                    continue
                if len(kids) == 1 and kids[0].type == "image":
                    t.attrSet("class", "fig")
                elif raw.endswith(":"):
                    t.attrSet("class", "lead")
            if t.type in ("fence", "code_block"):
                lang, title = fence_info(t.info or "")
                lines = to_mau(t.content, lang)
                body = "\n".join("".join(f'<span style="color:#{c}">{html.escape(s)}</span>' for c, s in ln) for ln in lines)
                head = f'<div class="ctitle">{html.escape(title)}</div>' if title else ""
                cls = "code" + (" titled" if title else "") + (" short" if len(lines) <= MA_NGAN else "")
                out.append(_html(f'{head}<pre class="{cls}">{body}</pre>\n'))
                i += 1
                continue
            if t.type == "html_block" and t.content.lstrip().startswith("<!--"):
                i += 1
                continue
            out.append(t)
            i += 1
        phan.append(md.renderer.render(out, md.options, {}))

    tpl = Document(TEMPLATE)
    logo = next(r.target_part.blob for r in tpl.part.rels.values() if r.reltype.endswith("/image"))
    logo_uri = "data:image/png;base64," + base64.b64encode(logo).decode()
    tv = [[c.text for c in row.cells] for row in tpl.tables[0].rows]
    tv_html = "".join("<tr>" + "".join((f"<th>{html.escape(c)}</th>" if k == 0 else f"<td style='text-align:{'left' if j == 1 else 'center'}'>{html.escape(c)}</td>")
                                       for j, c in enumerate(row)) + "</tr>" for k, row in enumerate(tv))
    bia = f"""<section class="cover front">
<p style="margin-top:6pt"><b style="font-size:14pt">ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH</b></p>
<p><b style="font-size:14pt;color:#1f3864">TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN</b></p>
<p style="font-size:14pt;margin-bottom:12pt">KHOA KHOA HỌC VÀ KỸ THUẬT THÔNG TIN</p>
<p><img src="{logo_uri}" style="width:3.35cm"></p>
<p style="margin-top:18pt"><b style="font-size:22pt;color:#1f3864">{html.escape(cfg['loai'])}</b></p>
<p><b style="font-size:16pt">QUẢN LÝ THÔNG TIN</b></p>
<p style="margin-bottom:15pt"><i style="font-size:14pt">MÃ LỚP: IE103.Q21.VB2</i></p>
<p style="margin-bottom:15pt"><b style="font-size:16pt;color:#2e75b6">{html.escape(cfg['tieu_de'])}</b></p>
<p>Giảng viên hướng dẫn: <b>TS. Võ Phương Bình</b></p><p>Nhóm sinh viên thực hiện: <b>Nhóm 1</b></p>
<table>{tv_html}</table>
<p style="margin-top:3cm"><i>{NGAY}</i></p></section>"""
    extra = ""
    if cfg["checklist"]:
        rows = "".join(f"<tr><td>{html.escape(x)}</td><td></td><td></td><td></td><td style='text-align:center'>X</td></tr>"
                       for x in cfg["checklist"])
        extra += (f"<section class='pb front'><p class='title'>CHECKLIST TIẾN ĐỘ HOÀN THÀNH</p><table><tr><th style='width:56%'>Nội dung</th>"
                  f"<th>25%</th><th>50%</th><th>75%</th><th>100%</th></tr>{rows}</table></section>")
    if cfg["nhan_xet"]:
        extra += "<section class='pb nx front'><p class='title'>NHẬN XÉT CỦA GIẢNG VIÊN</p><table>" + "<tr><td></td></tr>" * 18 + "</table></section>"
    ml = "".join(f'<li class="l{lv}"><a href="#{a}">{html.escape(tx)}</a></li>' for lv, a, tx in muc)
    dh = "".join(f'<li><a href="#{a}">{html.escape(tx)}</a></li>' for a, tx in hinh)
    db = "".join(f'<li><a href="#{a}">{html.escape(tx)}</a></li>' for a, tx in bang)
    vt = md.render(viet_tat_md("\n".join(src.read_text(encoding="utf-8") for src in cfg["src"]))).replace("<table>", "<table class='vt'>", 1)
    page = (f"<!doctype html><html lang='vi'><head><meta charset='utf-8'><title>{html.escape(cfg['tieu_de'])}</title>"
            f"<style>{CSS.replace('%HEADER%', cfg['header'])}{'' if cfg['ma_consolas'] else 'code { font: inherit; }'}</style>"
            f"</head><body>{bia}{extra}"
            f"<section class='pb toc front'><p class='title'>MỤC LỤC</p><ul>{ml}</ul></section>"
            f"<section class='pb toc front'><p class='title'>DANH MỤC HÌNH ẢNH</p><ul>{dh}</ul></section>"
            f"<section class='pb toc front'><p class='title'>DANH MỤC BẢNG</p><ul>{db}</ul></section>"
            f"<section class='pb front'><p class='title'>DANH MỤC TỪ VIẾT TẮT</p>{vt}</section>"
            + "".join(phan) + "</body></html>")
    out_pdf = cfg["out"].with_suffix(".pdf")
    with tempfile.TemporaryDirectory() as tmp:
        f = Path(tmp) / "doc.html"
        f.write_text(page, encoding="utf-8")
        subprocess.run([trinh_duyet(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        "--allow-file-access-from-files", f"--user-data-dir={tmp}/profile", f"--print-to-pdf={out_pdf}", f.as_uri()],
                       check=True, capture_output=True, timeout=300)
    print(f"✓ {out_pdf.relative_to(ROOT)}")


def _html(s):
    from markdown_it.token import Token
    t = Token("html_block", "", 0)
    t.content = s
    return t


def trinh_duyet() -> str:
    for c in (os.getenv("CHROME"), "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
              "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge", "/Applications/Chromium.app/Contents/MacOS/Chromium",
              shutil.which("google-chrome"), shutil.which("chromium"), shutil.which("msedge"),
              r"C:\Program Files\Google\Chrome\Application\chrome.exe", r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"):
        if c and Path(c).exists():
            return c
    raise SystemExit("Không tìm thấy Chrome / Edge / Chromium – đặt biến môi trường CHROME trỏ tới trình duyệt.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("tai_lieu", nargs="*", help=f"{', '.join(DOCS)} (mặc định: tất cả)")
    ap.add_argument("--pdf", action="store_true", help="xuất thêm PDF")
    a = ap.parse_args()
    if bad := [k for k in a.tai_lieu if k not in DOCS]:
        ap.error(f"không có tài liệu {bad}; chọn trong {list(DOCS)}")
    for k in a.tai_lieu or DOCS:
        build_docx(DOCS[k])
        if a.pdf:
            build_pdf(DOCS[k])
