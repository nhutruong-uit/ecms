"""Dựng slide thuyết trình đồ án → slides/ECMS_Nhom1.pptx (16:9, 17 slide cho buổi trình bày 15 phút).

    pip install -r scripts/requirements-docs.txt
    python scripts/build_slides.py

Thời lượng khoảng 14 phút, trong đó có 3 phút demo trực tiếp, còn 1 phút dự phòng. Mốc thời gian, người trình bày và ý
chính của từng slide nằm trong ghi chú (speaker notes); chân slide ghi người trình bày, góc trên ghi phần đang trình
bày và thanh tiến độ sáu phần.

Thiết kế: xanh than (logo UIT, báo cáo) làm màu chủ đạo, đỏ san hô (màu nút của ứng dụng) làm điểm nhấn, xanh lá cho
kết quả đạt; phông Calibri; biểu tượng Lucide (giấy phép ISC) đặt trong vòng tròn màu, lưu sẵn thành PNG trắng ở
slides/assets/icons/. Vẽ lại biểu tượng (cần rsvg-convert và gói npm lucide-static):

    python scripts/build_slides.py --icons node_modules/lucide-static/icons

Số liệu trên slide lấy từ báo cáo – sửa số liệu trong báo cáo thì sửa cả ở đây.
"""
import subprocess
import sys
from io import BytesIO
from math import cos, pi, sin
from pathlib import Path

from docx import Document
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "slides" / "ECMS_Nhom1.pptx"
ICON = ROOT / "slides" / "assets" / "icons"
CH2, CH4, DES = ROOT / "report" / "images" / "ch2", ROOT / "report" / "images" / "ch4", ROOT / "docs" / "design"
BIEU_TUONG = ("arrow-right", "boxes", "braces", "briefcase", "calculator", "calendar-check", "circle-check", "circle-x",
              "clipboard-pen", "code-xml", "container", "database", "eye-off", "file-check", "file-spreadsheet",
              "file-text", "fingerprint", "flag", "flask-conical", "gauge", "git-branch", "git-pull-request",
              "graduation-cap", "hard-drive", "history", "key-round", "layers", "layout-dashboard", "lightbulb",
              "list-tree", "lock", "network", "presentation", "receipt", "rocket", "rotate-ccw", "route", "scroll-text",
              "search", "server", "shield-check", "target", "triangle-alert", "users", "workflow", "zap")

# ----------------------------------------------------------------------------- bảng màu, kích thước (inch)
NAVY, NAVY2, BLUE, CORAL, GREEN, AMBER, VIOLET, SLATE = ("14284B", "1E3A6B", "2F5BD3", "EF5A4F", "1F9D6B", "D4891C",
                                                         "6D4FC2", "5B6B82")
INK, MUTED, LINE, CANVAS, WHITE, PALE, PALE2 = "1E293B", "5B6577", "E2E8F0", "F4F6FA", "FFFFFF", "C7D2E8", "8FA3C7"
NHAT = {NAVY: "E6EBF3", BLUE: "E9EFFC", CORAL: "FDEDEC", GREEN: "E5F5EE", AMBER: "FBF1E1", VIOLET: "EFEBFA",
        SLATE: "EDF0F4"}                                   # nền nhạt cùng tông
DAM = {NAVY: NAVY, BLUE: "2448A8", CORAL: "C2362C", GREEN: "157F55", AMBER: "9A5B07", VIOLET: "5136A3",
       SLATE: "3F4C61"}                                    # chữ đậm cùng tông, đặt trên nền nhạt
FONT = "Calibri"
W, H, MX = 13.333, 7.5, 0.6                                # khổ 16:9, lề trái / phải
CW, TOP, BOT = W - 2 * MX, 1.75, 6.75                      # vùng nội dung

# Sáu phần: tên, mô tả, người trình bày, biểu tượng, thời lượng (giây) – khớp tổng ghi chú các slide của phần
PHAN = [("Tổng quan", "Bài toán, hiện trạng, yêu cầu và giá trị của thông tin", "Như", "target", 100),
        ("Tổ chức thông tin", "ERD, lược đồ quan hệ, chuẩn hoá và cài đặt", "Trâm", "database", 110),
        ("Xử lý thông tin", "Truy vấn SQL và XML, an ninh, lập trình CSDL", "Dương · Như · Giang · Phú", "workflow", 165),
        ("Trình bày thông tin", "Menu theo vai trò, form, báo cáo và dashboard", "Giang", "layout-dashboard", 80),
        ("CSDL tiên tiến", "Hướng đối tượng, phân tán và NoSQL trên cùng bài toán", "Trâm · Phú · Dương", "network", 70),
        ("Kiểm thử, demo, kết luận", "Kiểm thử tích hợp, demo trực tiếp, hướng phát triển", "Như · cả nhóm (demo)", "flag",
         270)]
# Thành viên: tên gọi (dùng trong ghi chú), họ tên, MSSV
THANH_VIEN = [("Như", "Trương Quang Như", "25540022"), ("Trâm", "Đỗ Phạm Minh Trâm", "25540042"),
              ("Dương", "Đỗ Bình Dương", "25540008"), ("Phú", "Nguyễn Việt Phú", "25540025"),
              ("Giang", "Nguyễn Bảo Giang", "25540009")]
HO_TEN = {goi: ten for goi, ten, _ in THANH_VIEN}
TONG = 17

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(round(W * 914400)), Emu(round(H * 914400))
BLANK = prs.slide_layouts[6]
STT = {"n": 0, "giay": 0, "phan": [0] * 6, "icon": set()}
CAN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
NEO = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def I(v):
    return Emu(round(v * 914400))


def _rgb(h):
    return RGBColor.from_string(h)


def tron(a, b, t):
    """Màu nằm giữa a và b theo tỉ lệ t."""
    return "".join(f"{round(int(a[i:i + 2], 16) * (1 - t) + int(b[i:i + 2], 16) * t):02X}" for i in (0, 2, 4))


# ----------------------------------------------------------------------------- đo chữ để xếp dòng
_PHONG = {}


def _phong(dam):
    if dam not in _PHONG:
        ten = "calibrib.ttf" if dam else "calibri.ttf"
        thu_muc = [Path("/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts"), Path("C:/Windows/Fonts"),
                   Path.home() / "Library" / "Fonts", Path("/Library/Fonts")]
        tep = next((p for d in thu_muc if d.is_dir() for p in d.iterdir() if p.name.lower() == ten), None)
        _PHONG[dam] = ImageFont.truetype(str(tep), 100) if tep else None
    return _PHONG[dam]


def rong(chuoi, co, dam=False):
    """Bề rộng (inch) của chuỗi ở cỡ chữ co; không có phông Calibri thì ước lượng nửa em mỗi ký tự."""
    f = _phong(dam)
    return f.getlength(chuoi) * co / 7200 if f else len(chuoi) * co * 0.5 / 72


def so_dong(chuoi, co, w, dam=False):
    n = 0
    for doan in chuoi.replace("**", "").split("\n"):
        dong, n = "", n + 1
        for tu in doan.split(" "):
            thu = f"{dong} {tu}".strip()
            if dong and rong(thu, co, dam) > w:
                n, dong = n + 1, tu
            else:
                dong = thu
    return n


def cao(chuoi, co, w, dam=False, gian=1.0):
    return so_dong(chuoi, co, w, dam) * co * 1.2 * gian / 72


# ----------------------------------------------------------------------------- hình, chữ, ảnh
def _hieu_ung(sh, bong=False):
    """Ghi đè hiệu ứng của style mặc định (tránh bóng đổ ngoài ý muốn); bong=True: bóng mềm phía dưới."""
    eff = etree.SubElement(sh._element.spPr, qn("a:effectLst"))
    if bong:
        o = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(I(0.16)), dist=str(I(0.03)), dir="5400000",
                             algn="t", rotWithShape="0")
        etree.SubElement(etree.SubElement(o, qn("a:srgbClr"), val="0F172A"), qn("a:alpha"), val="12000")


def _to_mau(sh, nen, vien=None, day=0.75):
    if nen:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(nen)
    else:
        sh.fill.background()
    if vien:
        sh.line.color.rgb = _rgb(vien)
        sh.line.width = Pt(day)
    else:
        sh.line.fill.background()


def hinh(s, kieu, x, y, w, h, nen=None, vien=None, day=0.75, bo=None, bong=False):
    sh = s.shapes.add_shape(kieu, I(x), I(y), I(w), I(h))
    _to_mau(sh, nen, vien, day)
    if bo is not None:
        sh.adjustments[0] = min(0.5, bo / min(w, h))
    _hieu_ung(sh, bong)
    return sh


def _cat_nua(diem, truc, moc, duoi):
    """Cắt đa giác theo nửa mặt phẳng toạ độ[truc] ≤ moc (duoi=True) hoặc ≥ moc (Sutherland–Hodgman)."""
    trong = (lambda p: p[truc] <= moc) if duoi else (lambda p: p[truc] >= moc)
    ra = []
    for i, p in enumerate(diem):
        q = diem[i - 1]
        if trong(p) != trong(q):
            t = (moc - q[truc]) / (p[truc] - q[truc])
            ra.append((q[0] + t * (p[0] - q[0]), q[1] + t * (p[1] - q[1])))
        if trong(p):
            ra.append(p)
    return ra


def tron_trong_slide(s, cx, cy, r, nen):
    """Hình tròn tâm (cx, cy) bán kính r, bỏ phần nằm ngoài slide – không còn hình thừa ra mép khi chỉnh sửa."""
    diem = [(cx + r * cos(2 * pi * k / 360), cy + r * sin(2 * pi * k / 360)) for k in range(360)]
    for truc, moc, duoi in [(0, 0, False), (0, W, True), (1, 0, False), (1, H, True)]:
        diem = _cat_nua(diem, truc, moc, duoi)
    ff = s.shapes.build_freeform(I(diem[0][0]), I(diem[0][1]), scale=1.0)
    ff.add_line_segments([(I(x), I(y)) for x, y in diem[1:]])
    sh = ff.convert_to_shape()
    _to_mau(sh, nen)
    _hieu_ung(sh)
    return sh


def _dien_chu(tf, doan, co, mau, dam, can, neo, gian, sau, nhan, gian_chu, le):
    """doan: chuỗi hoặc danh sách; mỗi phần tử là chuỗi hoặc (chuỗi, tuỳ chọn riêng). '\\n' tách đoạn, **…** in đậm."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = I(le)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = NEO[neo]
    ds = []
    for d in doan if isinstance(doan, list) else [doan]:
        d, o = d if isinstance(d, tuple) else (d, {})
        ds += [(x, o) for x in d.split("\n")]
    for i, (d, o) in enumerate(ds):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = CAN[o.get("can", can)]
        p.line_spacing = o.get("gian", gian)
        p.space_after = Pt(o.get("sau", sau))
        nm = o.get("nhan", nhan)
        for k, seg in enumerate(d.split("**")):
            if not seg:
                continue
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name, f.size, f.bold = FONT, Pt(o.get("co", co)), o.get("dam", dam) or k % 2 == 1
            f.color.rgb = _rgb(nm if k % 2 and nm else o.get("mau", mau))
            rpr = r._r.get_or_add_rPr()
            rpr.set("lang", "vi-VN")
            if o.get("gian_chu", gian_chu):
                rpr.set("spc", str(int(o.get("gian_chu", gian_chu) * 100)))


def chu(s, x, y, w, h, doan, co=14, mau=INK, dam=False, can="l", neo="t", gian=1.0, sau=0, nhan=None, gian_chu=0, le=0):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    _dien_chu(tb.text_frame, doan, co, mau, dam, can, neo, gian, sau, nhan, gian_chu, le)
    return tb


def chu_hinh(sh, doan, co=12, mau=WHITE, dam=True, can="c", neo="m", gian=1.0, nhan=None, le=0.05):
    _dien_chu(sh.text_frame, doan, co, mau, dam, can, neo, gian, 0, nhan, 0, le)
    return sh


def the(s, x, y, w, h, nen=WHITE, vien=LINE, bong=True, bo=0.12):
    return hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, nen, vien, 0.75, bo, bong)


def nhan_tron(s, x, y, w, h, chuoi, nen, mau, co=12, dam=True):
    return chu_hinh(hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, nen, bo=h / 2), chuoi, co, mau, dam)


def hang_nhan(s, x, y, w, ds, mau, co=12, h=0.3, gian=0.08):
    """Hàng chip nối tiếp, hết chỗ thì xuống dòng; trả về y dưới hàng cuối."""
    px = x
    for chuoi in ds:
        pw = rong(chuoi, co, True) + 0.3
        if px > x and px + pw > x + w:
            px, y = x, y + h + 0.08
        nhan_tron(s, px, y, pw, h, chuoi, NHAT[mau], DAM[mau], co)
        px += pw + gian
    return y + h


def bieu_tuong(s, ten, x, y, d, nen):
    """Biểu tượng trắng trong vòng tròn màu – mô-típ chung của cả bộ slide."""
    STT["icon"].add(ten)
    hinh(s, MSO_SHAPE.OVAL, x, y, d, d, nen)
    k = 0.56
    s.shapes.add_picture(str(ICON / f"{ten}.png"), I(x + d * (1 - k) / 2), I(y + d * (1 - k) / 2), I(d * k), I(d * k))


def ti_le(path, cat=(0, 0, 0, 0)):
    with Image.open(path) as im:
        iw, ih = im.size
    l, t, r, b = cat
    return iw * (1 - l - r) / (ih * (1 - t - b))


def anh(s, path, x, y, w, h, cat=(0, 0, 0, 0), vien=None, bong=False):
    """Ảnh vừa khung w × h, giữ tỉ lệ, căn giữa; cat = phần cắt bỏ trái, trên, phải, dưới. Trả về vị trí thật."""
    ar = ti_le(path, cat)
    ww, hh = (w, w / ar) if w / ar <= h else (h * ar, h)
    x, y = x + (w - ww) / 2, y + (h - hh) / 2
    pic = s.shapes.add_picture(str(path), I(x), I(y), I(ww), I(hh))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cat
    if vien:
        pic.line.color.rgb = _rgb(vien)
        pic.line.width = Pt(0.75)
    _hieu_ung(pic, bong)
    return x, y, ww, hh


def trinh_duyet(s, path, x, y, w=None, h=None, cat=(0, 0, 0, 0)):
    """Ảnh chụp màn hình đặt trong khung cửa sổ trình duyệt; cho w thì tính h và ngược lại. Trả về (w, h) của khung."""
    thanh, le = 0.3, 0.06
    ar = ti_le(path, cat)
    if w is None:
        w = (h - thanh - le) * ar + 2 * le
    else:
        h = (w - 2 * le) / ar + thanh + le
    the(s, x, y, w, h, bo=0.1)
    for i, c in enumerate((CORAL, AMBER, GREEN)):
        hinh(s, MSO_SHAPE.OVAL, x + 0.16 + i * 0.17, y + 0.1, 0.1, 0.1, c)
    nhan_tron(s, x + 0.75, y + 0.065, min(2.6, w - 1.0), 0.17, "localhost:8501", CANVAS, MUTED, 9, False)
    pic = s.shapes.add_picture(str(path), I(x + le), I(y + thanh), I(w - 2 * le), I(h - thanh - le))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cat
    _hieu_ung(pic)
    return w, h


KHE_MUI_TEN = 0.56                                         # khe giữa hai thẻ có mũi tên: mũi tên 0.3 cách mỗi thẻ 0.13


def mui_ten(s, x, y, d=0.34, nen=SLATE):
    bieu_tuong(s, "arrow-right", x, y, d, nen)


def logo():
    """Phần hình của logo UIT trong mẫu báo cáo, bỏ dòng chữ bên dưới (cắt tại dải trắng đầu tiên). Trả về (ảnh, tỉ lệ)."""
    tpl = Document(ROOT / "report" / "template" / "mau-bao-cao-uit.docx")
    blob = next(r.target_part.blob for r in tpl.part.rels.values() if r.reltype.endswith("/image"))
    with Image.open(BytesIO(blob)) as im:
        im = im.convert("RGB")
        mau = im.convert("L").point(lambda v: 255 if v < 200 else 0)          # điểm có màu → trắng
        dau = mau.getbbox()[1]
        cuoi = next(y for y in range(dau, im.height) if not mau.crop((0, y, im.width, y + 1)).getbbox())
        im = im.crop(mau.crop((0, 0, im.width, cuoi)).getbbox())
        buf = BytesIO()
        im.save(buf, "PNG")
    return buf, im.width / im.height


# ----------------------------------------------------------------------------- khung slide
def slide_moi(nen):
    s = prs.slides.add_slide(BLANK)
    STT["n"] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = _rgb(nen)
    return s


def ghi_chu(s, nguoi, giay, y, phan=0):
    """Ghi chú người trình bày: mốc thời gian cộng dồn, người trình bày, các ý chính."""
    a = STT["giay"]
    STT["giay"] += giay
    if phan:
        STT["phan"][phan - 1] += giay
    t = lambda g: f"{g // 60}:{g % 60:02d}"  # noqa: E731
    s.notes_slide.notes_text_frame.text = (f"⏱ {t(a)} → {t(STT['giay'])} ({giay} giây) · Trình bày: {nguoi}\n"
                                           + "\n".join(f"• {x}" for x in y))


def dau_trang(s, phan, nhan, tieu_de, toi=False):
    """Số phần + nhãn phần, thanh tiến độ sáu phần, tiêu đề – cùng vị trí trên mọi slide nội dung."""
    if phan:
        chu_hinh(hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, 0.42, 0.5, 0.3, CORAL if toi else NAVY, bo=0.07),
                 f"{phan:02d}", 12)
        for i in range(6):
            c = CORAL if i == phan - 1 else ((PALE2 if toi else NAVY) if i < phan - 1 else ("2C4572" if toi else "D5DCE8"))
            hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, W - MX - 6 * 0.42 - 5 * 0.07 + i * 0.49, 0.53, 0.42, 0.08, c, bo=0.04)
    chu(s, MX + (0.64 if phan else 0), 0.42, 7, 0.3, nhan, 12, CORAL, True, neo="m", gian_chu=1.2)
    chu(s, MX, 0.8, CW, 0.62, tieu_de, 28, WHITE if toi else NAVY, True)


def ho_ten(nguoi):
    """'Như · Giang' → 'Trương Quang Như, Nguyễn Bảo Giang' – người xem không biết tên gọi trong nhóm."""
    return ", ".join(HO_TEN.get(t, t) for t in nguoi.split(" · "))


def chan_trang(s, nguoi, toi=False):
    m = PALE2 if toi else MUTED
    chu(s, MX, 7.0, 6, 0.25, "ECMS · Hệ thống quản lý trung tâm tiếng Anh · Nhóm 1 · IE103.Q21.VB2", 10, m)
    chu(s, W - MX - 6, 7.0, 6, 0.25, f"Trình bày: **{ho_ten(nguoi)}**   ·   {STT['n']:02d} / {TONG}", 10, m, can="r",
        nhan=WHITE if toi else NAVY)


def noi_dung(phan, tieu_de, nguoi, giay, y):
    s = slide_moi(CANVAS)
    dau_trang(s, phan, PHAN[phan - 1][0].upper(), tieu_de)
    chan_trang(s, nguoi)
    ghi_chu(s, nguoi, giay, y, phan)
    return s


def vong_trang_tri(s):
    """Hai vòng tròn mờ trên nền tối của slide mở đầu và kết thúc."""
    tron_trong_slide(s, 12.0, 1.5, 3.6, "1A3360")
    tron_trong_slide(s, 12.9, 7.0, 1.7, "1A3360")


def danh_sach(s, x, y, w, muc, ten_icon, nen, co=14, d=0.3, khoang=0.16):
    """Danh sách có dấu biểu tượng tròn nhỏ; trả về toạ độ y sau mục cuối."""
    for m in muc:
        h = cao(m, co, w - d - 0.18)
        bieu_tuong(s, ten_icon, x, y + max(0, (co * 1.2 / 72 - d) / 2), d, nen)
        chu(s, x + d + 0.18, y, w - d - 0.18, h + 0.05, m, co, INK, nhan=INK)
        y += max(h, d) + khoang
    return y


def tieu_de_the(s, x, y, w, ten_icon, nen, ten, phu=None, d=0.55):
    """Đầu thẻ: biểu tượng tròn + tên (+ dòng phụ)."""
    bieu_tuong(s, ten_icon, x, y, d, nen)
    doan = [(ten, {"co": 17, "dam": True, "mau": INK})] + ([(phu, {"co": 13, "mau": MUTED})] if phu else [])
    chu(s, x + d + 0.2, y - 0.05, w - d - 0.2, d + 0.1, doan, neo="m")


# ----------------------------------------------------------------------------- 1. bìa
def bia():
    s = slide_moi(NAVY)
    vong_trang_tri(s)
    x0 = 0.75
    anh_logo, ar = logo()
    lh = 0.7                                               # cao phần hình logo trong thẻ trắng
    the(s, x0, 0.62, lh * ar + 0.3, 0.95, WHITE, None, False, 0.14)
    s.shapes.add_picture(anh_logo, I(x0 + 0.15), I(0.62 + (0.95 - lh) / 2), height=I(lh))
    chu(s, x0 + lh * ar + 0.5, 0.62, 5.6, 0.95, [("ĐẠI HỌC QUỐC GIA THÀNH PHỐ HỒ CHÍ MINH", {"co": 11, "mau": PALE}),
                                                 ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN", {"co": 14, "dam": True}),
                                                 ("Khoa Khoa học và Kỹ thuật Thông tin", {"co": 11, "mau": PALE})],
        mau=WHITE, neo="m", sau=2)
    chu(s, x0, 2.3, 6.6, 0.32, "ĐỒ ÁN MÔN HỌC  ·  IE103 – QUẢN LÝ THÔNG TIN", 13, CORAL, True, gian_chu=1.5)
    chu(s, x0, 2.7, 6.8, 1.55, "Hệ thống quản lý\nTrung tâm Tiếng Anh", 44, WHITE, True, gian=0.92)
    nhan = dict(co=10, mau=PALE2, dam=True, gian_chu=1.2)
    chu(s, x0, 4.85, 2.8, 0.22, "GIẢNG VIÊN HƯỚNG DẪN", **nhan)
    chu(s, x0, 5.1, 2.8, 0.32, "TS. Võ Phương Bình", 15, WHITE, True)
    chu(s, x0, 5.62, 2.8, 0.22, "LỚP", **nhan)
    chu(s, x0, 5.87, 2.8, 0.32, "IE103.Q21.VB2 · Nhóm 1", 15, WHITE, True)
    chu(s, x0 + 2.95, 4.85, 3.8, 0.22, "THÀNH VIÊN", **nhan)
    for i, (_, ten, mssv) in enumerate(THANH_VIEN):
        chu(s, x0 + 2.95, 5.1 + i * 0.27, 3.0, 0.27, ten, 13, WHITE)
        chu(s, x0 + 5.95, 5.1 + i * 0.27, 0.85, 0.27, mssv, 12, PALE2, can="r")
    trinh_duyet(s, CH4 / "16_dashboard.jpg", 7.85, 1.05, w=4.88, cat=(0, 0, 0, 0.52))
    ghi_chu(s, "Như", 20, ["Chào thầy và các bạn; nhóm 1 trình bày đồ án ECMS – Hệ thống quản lý trung tâm tiếng Anh.",
                          "Giới thiệu nhanh năm thành viên; hình bên phải là dashboard thật của ứng dụng."])


# ----------------------------------------------------------------------------- 2. nội dung trình bày
def muc_luc():
    s = slide_moi(CANVAS)
    dau_trang(s, 0, "NỘI DUNG", "Bố cục trình bày")
    chan_trang(s, "Như")
    ghi_chu(s, "Như", 20, ["Sáu phần theo đúng quy trình quản lý thông tin của môn học; mỗi phần do người phụ trách "
                          "phần việc đó trình bày.", "Phần cuối có demo trực tiếp khoảng 3 phút."])
    xr, y0 = MX + 6.7, TOP + 0.35                         # vạch ngăn cột người trình bày, đầu dòng thứ nhất
    rh = (BOT - y0 - 5 * 0.1) / 6
    dau_cot = dict(co=10, mau=MUTED, dam=True, neo="m", gian_chu=1.2)
    chu(s, MX + 1.8, TOP, 4, 0.25, "PHẦN", **dau_cot)
    chu(s, xr + 0.3, TOP, 4, 0.25, "NGƯỜI TRÌNH BÀY", **dau_cot)
    for k, (ten, mo_ta, nguoi, ten_icon, _) in enumerate(PHAN):
        y = y0 + k * (rh + 0.1)
        the(s, MX, y, CW, rh)
        chu(s, MX + 0.3, y, 0.7, rh, f"{k + 1:02d}", 24, CORAL, True, neo="m")
        bieu_tuong(s, ten_icon, MX + 1.05, y + (rh - 0.48) / 2, 0.48, NAVY)
        chu(s, MX + 1.8, y, xr - MX - 2.0, rh, [(ten, {"co": 17, "dam": True, "mau": NAVY}), (mo_ta, {"co": 13, "mau": MUTED})],
            neo="m")
        hinh(s, MSO_SHAPE.RECTANGLE, xr, y + 0.15, 0.015, rh - 0.3, LINE)
        ds = ho_ten(nguoi).split(", ")
        dong = [", ".join(ds[i:i + 2]) for i in range(0, len(ds), 2)] if len(ds) > 3 else [", ".join(ds)]
        chu(s, xr + 0.3, y, MX + CW - xr - 0.55, rh, ",\n".join(dong), 14, INK, True, neo="m")   # bốn người: hai dòng


# ----------------------------------------------------------------------------- 3. bài toán
def bai_toan():
    s = noi_dung(1, "Bài toán: Ba chi nhánh, dữ liệu rời rạc trên Excel", "Như", 50, [
        "Chuỗi 3 chi nhánh, 4 chương trình (IELTS, TOEIC, Giao tiếp, Thiếu nhi); vòng đời học viên đi qua 7 bước.",
        "Hiện trạng: Mỗi chi nhánh một tệp Excel – dữ liệu trùng, không kiểm soát được sĩ số, trùng phòng; "
        "báo cáo doanh thu chép tay mất một đến hai ngày mỗi cuối tháng.",
        "Mục tiêu: Một CSDL tập trung, luật nghiệp vụ đặt ngay trong CSDL, ứng dụng web theo vai trò."])
    chu(s, MX, TOP, 6, 0.3, "QUY TRÌNH NGHIỆP VỤ BẢY BƯỚC", 11, MUTED, True, neo="m", gian_chu=1.2)
    buoc = ["Tư vấn", "Kiểm tra\nđầu vào", "Ghi danh,\nxếp lớp", "Thu\nhọc phí", "Học và\nđiểm danh", "Thi\ncuối khoá",
            "Báo cáo,\nthống kê"]
    bw, ov = (CW + 6 * 0.1) / 7, 0.1
    for i, b in enumerate(buoc):
        sh = hinh(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, MX + i * (bw - ov), TOP + 0.38, bw, 0.8,
                  tron(NAVY, BLUE, i / 6))
        sh.adjustments[0] = 0.3
        chu_hinh(sh, b, 12.5, le=0.02, gian=0.9)
    y, h, cw = TOP + 1.45, BOT - TOP - 1.45, (CW - 0.62) / 2
    cot = [(MX, CORAL, "file-spreadsheet", "Hiện trạng", "Mỗi chi nhánh một tệp Excel riêng", "circle-x", [
        "Hồ sơ học viên trùng lặp, sai lệch giữa các chi nhánh",
        "Không kiểm soát sĩ số, trùng phòng, khoá tiên quyết",
        "Báo cáo doanh thu chép tay, mất 1–2 ngày mỗi tháng",
        "Không phân quyền, không nhật ký, không sao lưu"]),
           (MX + cw + 0.62, GREEN, "target", "Mục tiêu của ECMS", "CSDL tập trung + ứng dụng web theo vai trò",
            "circle-check", ["Một nguồn dữ liệu cho 3 chi nhánh, 4 chương trình",
                             "Luật nghiệp vụ đặt tại CSDL: Ràng buộc, procedure, trigger",
                             "Báo cáo, dashboard tức thời thay cho tổng hợp tay",
                             "Phân quyền 6 tác nhân, nhật ký thay đổi, sao lưu định kỳ"])]
    for x, c, ic, ten, phu, dau, muc in cot:
        the(s, x, y, cw, h)
        tieu_de_the(s, x + 0.3, y + 0.28, cw - 0.6, ic, c, ten, phu)
        danh_sach(s, x + 0.3, y + 1.2, cw - 0.6, muc, dau, c, 14, khoang=0.24)
    mui_ten(s, MX + cw + 0.31 - 0.25, y + h / 2 - 0.25, 0.5, NAVY)


# ----------------------------------------------------------------------------- 4. yêu cầu
def yeu_cau():
    s = noi_dung(1, "Yêu cầu: Sáu tác nhân và giá trị của thông tin", "Như", 50, [
        "6 tác nhân, mỗi tác nhân ánh xạ sang một role trong SQL Server; 22 yêu cầu chức năng, 8 phi chức năng, 7 báo cáo.",
        "DIKW trên số liệu thật: 10.063 lượt điểm danh → HV0021 chuyên cần 56,25% → nhóm chuyên cần dưới 80% có tỷ lệ "
        "không đạt 33,3%, gấp 3 nhóm còn lại (11,1%) → hành động: Gọi điện, xếp học bù sớm."])
    sw = (CW - 0.9) / 4
    for i, (so, nhan, phu) in enumerate([("6", "tác nhân", "mỗi tác nhân một role CSDL"),
                                         ("22", "yêu cầu chức năng", "FR01 – FR22"),
                                         ("8", "phi chức năng", "toàn vẹn, bảo mật, truy vết…"),
                                         ("7", "báo cáo", "xuất PDF, Excel, Word")]):
        x = MX + i * (sw + 0.3)
        the(s, x, TOP, sw, 1.0)
        chu(s, x + 0.25, TOP, 0.95, 1.0, so, 36, CORAL if i == 0 else NAVY, True, neo="m")
        nx = x + 0.25 + rong(so, 36, True) + 0.2
        chu(s, nx, TOP, x + sw - nx - 0.15, 1.0, [(nhan, {"co": 15, "dam": True}), (phu, {"co": 12, "mau": MUTED})],
            neo="m")
    y0, lw = TOP + 1.25, 6.9
    mw, mh = (lw - 0.25) / 2, (BOT - y0 - 0.4) / 3
    vai_tro = [("Quản lý", "briefcase", NAVY, "Danh mục, báo cáo, dashboard, nhật ký"),
               ("Giáo vụ", "clipboard-pen", BLUE, "Tư vấn, mở lớp, xếp lịch, ghi danh"),
               ("Kế toán", "calculator", GREEN, "Thu học phí, công nợ, bảng lương"),
               ("Giáo viên", "presentation", CORAL, "Điểm danh, nhập điểm, cảnh báo chuyên cần"),
               ("Học viên", "graduation-cap", VIOLET, "Tra cứu lịch học, điểm, học phí của mình"),
               ("Quản trị hệ thống", "server", SLATE, "Tài khoản, phân quyền, sao lưu")]
    for k, (ten, ic, c, mo_ta) in enumerate(vai_tro):
        x, y = MX + k % 2 * (mw + 0.25), y0 + k // 2 * (mh + 0.2)
        the(s, x, y, mw, mh)
        bieu_tuong(s, ic, x + 0.22, y + (mh - 0.5) / 2, 0.5, c)
        chu(s, x + 0.9, y + 0.12, mw - 1.05, mh - 0.24, [(ten, {"co": 15, "dam": True}), (mo_ta, {"co": 13, "mau": MUTED})],
            neo="m")
    x, w, h = MX + lw + 0.3, CW - lw - 0.3, BOT - y0
    the(s, x, y0, w, h)
    chu(s, x + 0.3, y0 + 0.2, w - 0.6, 0.32, "Từ dữ liệu đến hành động (DIKW)", 15, INK, True)
    bac = [("DỮ LIỆU", "10.063 lượt điểm danh", SLATE), ("THÔNG TIN", "HV0021: Chuyên cần 56,25%", BLUE),
           ("TRI THỨC", "Chuyên cần < 80%: Trượt gấp 3 lần", NAVY), ("TRÍ TUỆ", "Gọi điện, xếp học bù sớm", CORAL)]
    bw, bh = 3.3, 0.66
    buoc_x = (w - 0.5 - bw) / 3
    buoc_y = (h - 0.75 - 0.2 - bh) / 3
    for i, (ten, gt, c) in enumerate(bac):
        bx, by = x + 0.25 + i * buoc_x, y0 + h - 0.2 - bh - i * buoc_y
        chu_hinh(hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, bh, c, bo=0.08),
                 [(ten, {"co": 10, "mau": "DCE3F0", "gian_chu": 1.2}), (gt, {"co": 13.5})], can="l", le=0.18)


# ----------------------------------------------------------------------------- 5. ERD
def erd():
    s = noi_dung(2, "Mô hình khái niệm: ERD 14 thực thể, 22 mối kết hợp", "Trâm", 60, [
        "ERD ký hiệu Chen: 14 tập thực thể, 22 mối kết hợp, bản số (min, max) đặt sát thực thể tham gia.",
        "Chỉ vào: Thực thể yếu PHONGHOC, BUOIHOC (viền đôi); ISA NHANVIEN → GIAOVIEN; LỊCH HỌC ba ngôi.",
        "Sơ đồ lớp UML vẽ cho cùng bài toán để đối chiếu: Kế thừa, thành phần, kết tập."])
    lw = 8.5
    the(s, MX, TOP, lw, BOT - TOP)
    anh(s, DES / "erd.png", MX + 0.15, TOP + 0.15, lw - 0.3, BOT - TOP - 0.3)
    x, w = MX + lw + 0.35, CW - lw - 0.35
    y = TOP + 0.05
    for ic, c, ten, mo_ta in [("layers", BLUE, "Thực thể yếu", "PHONGHOC, BUOIHOC mượn khoá của CHINHANH, LOPHOC"),
                              ("git-branch", GREEN, "Chuyên biệt hoá", "NHANVIEN → GIAOVIEN, rời nhau, không toàn phần"),
                              ("network", CORAL, "Liên kết ba ngôi", "LỊCH HỌC nối lớp – ca học – phòng học"),
                              ("list-tree", VIOLET, "Đủ loại thuộc tính", "Khoá, phức hợp, đa trị (chứng chỉ), suy diễn"),
                              ("boxes", NAVY, "Sơ đồ lớp UML", "Kế thừa, thành phần, kết tập; đối chiếu với ERD")]:
        bieu_tuong(s, ic, x, y, 0.5, c)
        chu(s, x + 0.68, y - 0.04, w - 0.68, 0.95, [(ten, {"co": 15, "dam": True, "mau": INK, "sau": 1}),
                                                    (mo_ta, {"co": 13, "mau": MUTED})])
        y += 1.0


# ----------------------------------------------------------------------------- 6. lược đồ và cài đặt
def luoc_do():
    s = noi_dung(2, "Mô hình logic và cài đặt: 22 bảng trên SQL Server", "Trâm", 50, [
        "Chuyển ERD sang 22 lược đồ quan hệ, kiểm tra dạng chuẩn trên khoá dự tuyển: 16 BCNF, 4 đạt 2NF do lưu giá trị "
        "suy diễn, 2 vi phạm 2NF có chủ đích để khai báo khoá ngoại phức hợp thay cho trigger.",
        "Cài đặt: 33 khoá ngoại, 53 CHECK, 8 cột tính toán, 23 chỉ mục; 21 ràng buộc toàn vẹn có bảng tầm ảnh hưởng.",
        "Dữ liệu mẫu sinh theo trình tự nghiệp vụ, có sẵn trường hợp biên."])
    lw = 6.95
    the(s, MX, TOP, lw, BOT - TOP)
    anh(s, CH2 / "db-diagram-1.png", MX + 0.15, TOP + 0.2, lw - 0.3, BOT - TOP - 0.85)
    chu(s, MX + 0.15, BOT - 0.5, lw - 0.3, 0.3, "Sơ đồ CSDL nhóm đào tạo – nhân sự (13 / 22 bảng)", 12, MUTED, can="c",
        neo="m")
    x, w = MX + lw + 0.3, CW - lw - 0.3
    tw = (w - 0.4) / 3
    for k, (so, nhan) in enumerate([("22", "bảng"), ("33", "khoá ngoại"), ("53", "ràng buộc CHECK"),
                                    ("23", "chỉ mục bổ sung"), ("21", "RBTV đặc tả"), ("8", "cột tính toán")]):
        tx, ty = x + k % 3 * (tw + 0.2), TOP + k // 3 * 1.2
        the(s, tx, ty, tw, 1.0)
        chu(s, tx, ty + 0.1, tw, 0.5, so, 28, NAVY, True, can="c", neo="m")
        chu(s, tx, ty + 0.6, tw, 0.3, nhan, 12, MUTED, can="c")
    y = TOP + 2.4
    the(s, x, y, w, 1.65)
    chu(s, x + 0.25, y + 0.18, w - 0.5, 0.3, "Dạng chuẩn trên khoá dự tuyển", 14, INK, True)
    bx, bw = x + 0.25, w - 0.5
    for so, nhan, c in [(16, "16 · BCNF", NAVY), (4, "4 · 2NF", BLUE), (2, "2", CORAL)]:
        sw = bw * so / 22
        chu_hinh(hinh(s, MSO_SHAPE.RECTANGLE, bx, y + 0.55, sw, 0.42, c), nhan, 13)
        bx += sw
    chu(s, x + 0.25, y + 1.05, w - 0.5, 0.5, "2NF: Lưu giá trị suy diễn để tra nhanh · **2 bảng vi phạm 2NF có chủ đích** "
        "để khai báo khoá ngoại phức hợp", 12, MUTED, nhan=DAM[CORAL])
    y += 1.85
    the(s, x, y, w, BOT - y, NHAT[BLUE], None, False)
    bieu_tuong(s, "database", x + 0.2, y + (BOT - y - 0.42) / 2, 0.42, BLUE)
    chu(s, x + 0.78, y, w - 0.95, BOT - y, "**Dữ liệu mẫu theo nghiệp vụ:** 349 học viên · 25 lớp · 539 đăng ký · "
        "10.063 lượt điểm danh", 13, INK, neo="m", nhan=DAM[BLUE])


# ----------------------------------------------------------------------------- 7. truy vấn
def truy_van():
    s = noi_dung(3, "Truy vấn: 31 câu SQL, 18 câu XML, đo bằng số trang đọc", "Dương", 55, [
        "31 truy vấn SQL chia 6 nhóm kỹ thuật, mỗi câu trả lời một câu hỏi quản lý thực tế (gợi ý chạy thử Q18, Q24).",
        "Đo trên bảng điểm danh nhân bản 1.000.000 dòng: Chỉ mục giảm từ 6.023 trang xuống 3 trang (một đăng ký) và "
        "150 trang (một lớp).",
        "XML: Kết quả kiểm tra đầu vào lưu cột XML có XML Schema, điểm ngoài 0–10 bị từ chối; chỉ mục XML có lợi ở quy mô "
        "thật nhưng bất lợi khi gấp 30 lần – không phải chỉ mục nào cũng tốt."])
    lw = 6.6
    chu(s, MX, TOP, lw, 0.3, "31 TRUY VẤN SQL THEO 6 NHÓM KỸ THUẬT", 11, MUTED, True, neo="m", gian_chu=1.2)
    tw, th = (lw - 0.25) / 2, (BOT - TOP - 0.45 - 0.44) / 3
    for k, (so, ten, kt) in enumerate([("5", "Cơ bản", "LIKE, BETWEEN, TOP … WITH TIES"),
                                       ("5", "Kết nối", "INNER, LEFT, RIGHT, FULL JOIN, tự kết nối"),
                                       ("5", "Gom nhóm", "GROUP BY, HAVING, ROLLUP"),
                                       ("5", "Truy vấn lồng", "IN, NOT EXISTS, tương quan, ALL"),
                                       ("5", "Tập hợp", "UNION, INTERSECT, EXCEPT, phép chia"),
                                       ("6", "Nâng cao", "CTE đệ quy, hàm cửa sổ, PIVOT")]):
        x, y = MX + k % 2 * (tw + 0.25), TOP + 0.45 + k // 2 * (th + 0.22)
        the(s, x, y, tw, th)
        chu_hinh(hinh(s, MSO_SHAPE.OVAL, x + 0.22, y + (th - 0.55) / 2, 0.55, 0.55, NAVY), so, 18)
        chu(s, x + 0.95, y + 0.1, tw - 1.1, th - 0.2, [(ten, {"co": 16, "dam": True, "sau": 2}),
                                                      (kt, {"co": 13, "mau": MUTED})], neo="m")
    x, w = MX + lw + 0.3, CW - lw - 0.3
    h1 = 2.55
    the(s, x, TOP, w, h1)
    tieu_de_the(s, x + 0.3, TOP + 0.25, w - 0.6, "zap", CORAL, "Chỉ mục trên 1.000.000 dòng điểm danh", d=0.5)
    for i, (mo_ta, sau, giam) in enumerate([("Chuyên cần của một đăng ký", "3 trang", "−99,9%"),
                                            ("Thống kê một lớp", "150 trang", "−97,5%")]):
        y = TOP + 0.9 + i * 0.75
        chu(s, x + 0.3, y, w - 0.6, 0.25, mo_ta, 13, MUTED)
        chu(s, x + 0.3, y + 0.25, 3.6, 0.5, f"6.023  →  **{sau}**", 26, SLATE, nhan=CORAL, neo="m")
        nhan_tron(s, x + w - 1.3, y + 0.33, 1.0, 0.34, giam, NHAT[GREEN], DAM[GREEN], 13)
    y = TOP + h1 + 0.25
    the(s, x, y, w, BOT - y)
    tieu_de_the(s, x + 0.3, y + 0.22, w - 0.6, "code-xml", BLUE, "XML có XML Schema: 18 truy vấn", d=0.5)
    yy = hang_nhan(s, x + 0.3, y + 0.8, w - 0.6, ("query()", "value()", "exist()", "nodes()", "modify()", "FOR XML",
                                                   "OPENXML"), BLUE)
    chu(s, x + 0.3, yy + 0.08, w - 0.6, 0.5, "Chỉ mục XML có lợi ở 357 bài (1.969 → 707 trang) nhưng **bất lợi ở 10.710 bài** "
        "(1.956 → 40.490 trang)", 13, INK, nhan=DAM[CORAL])


# ----------------------------------------------------------------------------- 8. an ninh
def an_ninh():
    s = noi_dung(3, "An ninh dữ liệu: Bảo vệ theo 7 lớp, phục hồi dưới 1 giây", "Như · Giang", 60, [
        "Như (30 giây): Đi từ ngoài vào trong – TLS, băm lặp có salt và tạm khoá, login ứng dụng bị DENY mọi bảng, danh "
        "tính ghi vào SESSION_CONTEXT chỉ đọc nên mã bị chèn không đổi được; thiếu ngữ cảnh thì từ chối.",
        "Giang (30 giây): View bảo mật theo cột, mã hoá lưu trữ, kiểm toán; sao lưu Full tuần, vi sai ngày, nhật ký 15 "
        "phút → RPO 15 phút; diễn tập phục hồi đúng thời điểm mất 607–810 ms."])
    lw = 7.25
    lop = [("lock", "Mã hoá đường truyền", "TLS giữa ứng dụng và SQL Server"),
           ("fingerprint", "Xác thực", "Băm lặp 10.000 vòng SHA2_256 có salt; sai 5 lần tạm khoá 15 phút"),
           ("key-round", "Quyền tối thiểu", "6 role; login ứng dụng bị DENY mọi bảng, chỉ gọi procedure, đọc view"),
           ("users", "Phân quyền theo dòng", "SESSION_CONTEXT chỉ đọc; thiếu ngữ cảnh thì từ chối"),
           ("eye-off", "Che, ẩn theo cột", "View bảo mật ẩn lương; Dynamic Data Masking số điện thoại"),
           ("hard-drive", "Mã hoá lưu trữ", "TDE, mã hoá cột, bản sao lưu mã hoá AES-256"),
           ("scroll-text", "Kiểm toán", "Trigger nhật ký JSON, SQL Server Audit, bảng ledger")]
    rh, gap = 0.62, (BOT - TOP - 7 * 0.62) / 6
    for i, (ic, ten, cc) in enumerate(lop):
        y = TOP + i * (rh + gap)
        the(s, MX, y, lw, rh, bong=False)
        bieu_tuong(s, ic, MX + 0.12, y + 0.1, 0.42, tron(BLUE, NAVY, i / 6))
        chu(s, MX + 0.7, y, 2.1, rh, [(f"LỚP {i + 1}", {"co": 9, "mau": MUTED, "dam": True, "gian_chu": 1}),
                                     (ten, {"co": 14, "dam": True})], neo="m")
        chu(s, MX + 2.85, y, lw - 3.0, rh, cc, 13, MUTED, neo="m")
    x, w = MX + lw + 0.3, CW - lw - 0.3
    h1 = 3.25
    the(s, x, TOP, w, h1)
    tieu_de_the(s, x + 0.3, TOP + 0.25, w - 0.6, "history", GREEN, "Sao lưu và phục hồi", d=0.5)
    cw = (w - 0.6 - 0.2) / 2
    for i, (so, nhan) in enumerate([("15 phút", "RPO – mất tối đa"), ("< 1 giây", "RTO đo được 607–810 ms")]):
        cx = x + 0.3 + i * (cw + 0.2)
        the(s, cx, TOP + 0.95, cw, 1.1, NHAT[NAVY], None, False, 0.1)
        chu(s, cx, TOP + 1.02, cw, 0.55, so, 26, NAVY, True, can="c", neo="m")
        chu(s, cx, TOP + 1.57, cw, 0.4, nhan, 12, MUTED, can="c")
    chu(s, x + 0.3, TOP + 2.25, w - 0.6, 0.9, ["Full hằng tuần · vi sai hằng ngày · nhật ký 15 phút",
                                              "Phục hồi đúng thời điểm (STOPAT), DBCC CHECKDB không lỗi"], 13, INK, sau=4)
    y = TOP + h1 + 0.25
    the(s, x, y, w, BOT - y, NHAT[GREEN], None, False)
    bieu_tuong(s, "shield-check", x + 0.3, y + (BOT - y - 0.6) / 2, 0.6, GREEN)
    chu(s, x + 1.1, y, 1.0, BOT - y, "53", 36, DAM[GREEN], True, neo="m")
    chu(s, x + 1.95, y, w - 2.15, BOT - y, "ca kiểm thử phân quyền đều đạt – vai trò bị chặn, phạm vi chi nhánh, "
        "thiếu ngữ cảnh", 13, INK, neo="m")


# ----------------------------------------------------------------------------- 9. lập trình CSDL
def lap_trinh():
    s = noi_dung(3, "Lập trình CSDL: Luật nghiệp vụ nằm trọn trong CSDL", "Phú", 50, [
        "40 procedure, 17 function, 14 trigger và cursor đóng gói mọi quy tắc nghiệp vụ; giao diện chỉ gọi procedure.",
        "Ví dụ ghi danh: sp_DangKyHoc khoá dòng lớp bằng UPDLOCK, HOLDLOCK nên hai giáo vụ ghi danh cùng lúc không "
        "vượt sĩ số; lỗi nghiệp vụ trả bằng THROW 50004 kèm thông báo tiếng Việt, giao diện hiện nguyên văn.",
        "Cursor được so sánh công bằng với cách xử lý theo tập hợp: Tập hợp nhanh hơn 1,1–6,9 lần."])
    cw = (CW - 0.9) / 4
    for i, (so, ten, ic, c, vd, ket_qua) in enumerate([
            ("40", "Stored procedure", "workflow", NAVY, "sp_DangKyHoc, sp_ThuHocPhi nhận chi tiết dạng JSON",
             "39 ca kiểm thử đạt"),
            ("17", "Function", "braces", BLUE, "fn_GoiYKhoaHoc, fn_BangDiemLop, fn_NguoiDungHienTai",
             "Dùng trong view, CHECK"),
            ("14", "Trigger", "zap", CORAL, "Sĩ số, tổng tiền phiếu thu, chặn trùng phòng, nhật ký", "22 ca kiểm thử đạt"),
            ("3", "Cursor", "rotate-ccw", VIOLET, "Sinh buổi học bỏ ngày lễ, tính lương, cảnh báo",
             "Tập hợp nhanh hơn 1,1–6,9×")]):
        x = MX + i * (cw + 0.3)
        the(s, x, TOP, cw, 2.38)
        chu(s, x + 0.25, TOP + 0.15, 1.4, 0.65, so, 38, c if c != NAVY else NAVY, True, neo="m")
        bieu_tuong(s, ic, x + cw - 0.75, TOP + 0.22, 0.5, c)
        chu(s, x + 0.25, TOP + 0.82, cw - 0.5, 0.3, ten, 16, INK, True)
        chu(s, x + 0.25, TOP + 1.13, cw - 0.5, 0.66, vd, 12.5, MUTED)
        nhan_tron(s, x + 0.25, TOP + 1.86, cw - 0.5, 0.32, ket_qua, NHAT[c if c != NAVY else GREEN],
                  DAM[c if c != NAVY else GREEN], 12)
    y = TOP + 2.6
    the(s, MX, y, CW, BOT - y)
    chu(s, MX + 0.3, y + 0.2, CW - 0.6, 0.3, "Một lượt ghi danh đi qua CSDL", 15, INK, True)
    nut = [("Giao diện", "Bấm “Ghi danh”, gọi procedure", SLATE),
           ("sp_DangKyHoc", "Khoá dòng lớp (UPDLOCK); kiểm tra sĩ số, tiên quyết", NAVY),
           ("Trigger", "Cập nhật sĩ số, ghi nhật ký thay đổi", CORAL),
           ("COMMIT / THROW", "Hợp lệ thì lưu; vi phạm thì THROW 50004", BLUE),
           ("Giao diện", "Hiện nguyên văn “Lớp đã đủ sĩ số.”", SLATE)]
    ag = KHE_MUI_TEN
    nw = (CW - 0.6 - 4 * ag) / 5
    for i, (ten, mo_ta, c) in enumerate(nut):
        nx = MX + 0.3 + i * (nw + ag)
        the(s, nx, y + 0.65, nw, 1.1, NHAT[c], None, False, 0.1)
        chu(s, nx + 0.15, y + 0.65, nw - 0.3, 1.1, [(ten, {"co": 13.5, "dam": True, "mau": DAM[c], "sau": 2}),
                                                     (mo_ta, {"co": 12, "mau": INK})], neo="m")
        if i < 4:
            mui_ten(s, nx + nw + (ag - 0.3) / 2, y + 1.05, 0.3)
    hang_nhan(s, MX + 0.3, y + 1.95, CW - 0.6, ("XACT_ABORT + savepoint", "sp_getapplock, UPDLOCK chống tranh chấp",
                                                "Mã yêu cầu chống lập phiếu thu trùng"), NAVY, h=0.32, gian=0.15)


# ----------------------------------------------------------------------------- 10. menu và form
def giao_dien():
    s = noi_dung(4, "Giao diện: Menu theo vai trò, form kiểm tra hai lớp", "Giang", 40, [
        "Menu rộng và nông: 2 cấp, 8 nhóm theo nghiệp vụ, mục bị ẩn theo vai trò và chặn cả khi gõ thẳng URL "
        "(ảnh: Giáo viên chỉ thấy Học tập, Báo cáo, Hệ thống).",
        "Form master–detail phiếu thu: Một phiếu nhiều khoản, thu một phần, lưu trong một giao dịch sp_ThuHocPhi.",
        "Kiểm tra hai lớp: Lỗi hiện ngay dưới ô nhập, CSDL kiểm tra lại lần nữa."])
    fh = 3.87
    w1, _ = trinh_duyet(s, CH4 / "02b_trang_chu_giao_vien.jpg", MX, TOP, h=fh)
    w2 = CW - w1 - 0.3
    ar = (w2 - 0.12) / (fh - 0.36)                        # tỉ lệ cần cho ảnh phiếu thu → cắt bớt trên, dưới
    with Image.open(CH4 / "14_phieu_thu.jpg") as im:
        bo = 1 - im.size[0] / ar / im.size[1]
    trinh_duyet(s, CH4 / "14_phieu_thu.jpg", MX + w1 + 0.3, TOP, w=w2, cat=(0, 0.167, 0, bo - 0.167))
    for x, w, ic, c, cap in [(MX, w1, "list-tree", NAVY, "**Menu 2 cấp, 8 nhóm theo nghiệp vụ** – ẩn mục theo vai trò, "
                                                         "chặn cả truy cập thẳng bằng URL"),
                             (MX + w1 + 0.3, w2, "receipt", CORAL, "**Form master–detail** – một phiếu thu nhiều khoản, "
                                                                     "lưu trong một giao dịch")]:
        bieu_tuong(s, ic, x, TOP + fh + 0.2, 0.4, c)
        chu(s, x + 0.55, TOP + fh + 0.15, w - 0.55, 0.5, cap, 13, MUTED, nhan=INK, neo="m")
    pw = (CW - 0.6) / 4
    for i, y_chinh in enumerate(["Lỗi hiện ngay dưới ô nhập", "CSDL kiểm tra lại lần hai", "Xoá phải xác nhận",
                                 "Form chỉ gọi view, procedure"]):
        nhan_tron(s, MX + i * (pw + 0.2), BOT - 0.34, pw, 0.34, y_chinh, NHAT[NAVY], NAVY, 12.5)


# ----------------------------------------------------------------------------- 11. báo cáo, dashboard
def bao_cao():
    s = noi_dung(4, "Báo cáo và dashboard: Từ dữ liệu đến quyết định", "Giang", 40, [
        "7 báo cáo: Định kỳ, theo yêu cầu, tổng hợp, biểu mẫu; xuất PDF (A4, phiếu thu A5), Excel, Word, giữ tô màu có "
        "điều kiện và dòng trung bình, nhỏ nhất, lớn nhất.",
        "Báo cáo PDF có đủ Report / Page / Group header – footer, đánh số Trang X / Y – đối chiếu Crystal Report (Lab 4).",
        "Dashboard 6 KPI có mốc so sánh (kỳ trước, cùng kỳ, mục tiêu) – tầng Tri thức của DIKW."])
    w1, _ = trinh_duyet(s, CH4 / "16_dashboard.jpg", MX, TOP, h=BOT - TOP, cat=(0, 0, 0, 0.52))
    x = MX + w1 + 0.35
    pw = (BOT - TOP - 0.45) * ti_le(CH4 / "20_pdf_doanh_thu.jpg")
    anh(s, CH4 / "20_pdf_doanh_thu.jpg", x, TOP, pw, BOT - TOP - 0.45, vien=LINE, bong=True)
    chu(s, x, BOT - 0.3, pw, 0.3, "Báo cáo doanh thu xuất PDF", 12, MUTED, can="c", neo="m")
    x += pw + 0.35
    w = MX + CW - x
    bh = (BOT - TOP - 0.4) / 3
    for i, (ic, c, so, nhan, mo_ta) in enumerate([
            ("file-text", NAVY, "7", "báo cáo", "Định kỳ, theo yêu cầu, tổng hợp, biểu mẫu"),
            ("file-check", BLUE, "3", "định dạng xuất", "PDF (A4, A5), Excel, Word; giữ tô màu, dòng TB / Min / Max"),
            ("gauge", CORAL, "6", "KPI dashboard", "So với kỳ trước, cùng kỳ, mục tiêu; doanh thu đọc từ indexed view")]):
        y = TOP + i * (bh + 0.2)
        the(s, x, y, w, bh)
        bieu_tuong(s, ic, x + 0.25, y + 0.25, 0.5, c)
        chu(s, x + 0.9, y + 0.15, 0.6, 0.7, so, 30, c, True, neo="m")
        chu(s, x + 0.9 + rong(so, 30, True) + 0.12, y + 0.15, w - 1.3, 0.7, nhan, 16, INK, True, neo="m")
        chu(s, x + 0.25, y + 0.9, w - 0.5, bh - 1.0, mo_ta, 13, MUTED)


# ----------------------------------------------------------------------------- 12. CSDL tiên tiến
def tien_tien():
    s = noi_dung(5, "Bốn mô hình CSDL: Quan hệ làm lõi, các mô hình khác bổ trợ", "Trâm · Phú · Dương", 70, [
        "Trâm (20 giây): Hướng đối tượng – lược đồ ODL 16 lớp, OID, ba cách ánh xạ kế thừa; nạp lười sinh 889 câu SQL "
        "so với 1 câu khi nạp trước.",
        "Phú (25 giây): Phân tán theo chi nhánh – 3 trạm, phân mảnh ngang/dọc/hỗn hợp, nửa kết giảm 73% byte truyền, "
        "cam kết hai pha qua MSDTC chạy thật giữa hai máy.",
        "Dương (25 giây): NoSQL – 14 bảng thành 4 collection MongoDB; đọc trọn hồ sơ 0,70 ms so với 3,62 ms nhưng sửa "
        "dữ liệu dùng chung phải sửa mọi bản nhúng. Kết luận: Polyglot persistence, SQL Server là nguồn gốc."])
    c0 = 2.45
    cw = (CW - c0) / 4
    mo_hinh = [("Quan hệ", "database", NAVY, "SQL Server, đã cài đặt"),
               ("Hướng đối tượng", "boxes", BLUE, "ODL 16 lớp, OID"),
               ("Phân tán", "network", GREEN, "3 trạm, cam kết 2 pha"),
               ("NoSQL", "braces", CORAL, "MongoDB, 4 collection")]
    hh, rh, dh = 0.85, 0.54, 0.72
    yb = TOP + hh + 0.1                                    # đầu bảng
    the(s, MX, yb, CW, 5 * rh + dh, bong=True)
    hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX + c0 + 0.04, TOP, cw - 0.08, hh + 0.1 + 5 * rh + dh, NHAT[NAVY],
         bo=0.12)                                          # cột được chọn
    for j, (ten, ic, c, phu) in enumerate(mo_hinh):
        x = MX + c0 + j * cw
        bieu_tuong(s, ic, x + 0.15, TOP + 0.12, 0.46, c)
        chu(s, x + 0.7, TOP + 0.1, cw - 0.8, 0.5, [(ten, {"co": 15, "dam": True, "mau": INK}),
                                                   (phu, {"co": 12, "mau": MUTED})], neo="m")
    hang = [("Mô hình hoá bài toán", [(2, "Tách bảng đa trị"), (3, "Tự nhiên nhất"), (2, "Như quan hệ"),
                                      (2, "Tự nhiên cho hồ sơ")]),
            ("Toàn vẹn, giao dịch", [(3, "ACID, FK, CHECK"), (1, "Tự viết trong lớp"), (2, "ACID cục bộ + 2PC"),
                                     (1, "Không khoá ngoại")]),
            ("Thống kê, báo cáo", [(3, "Mạnh nhất"), (1, "Vấn đề N+1"), (2, "Liên trạm tốn kém"), (2, "Aggregation")]),
            ("Đọc trọn một hồ sơ", [(1, "Nhiều phép JOIN"), (2, "Duyệt đối tượng"), (1, "Như quan hệ"),
                                    (3, "Một document")]),
            ("Mở rộng quy mô", [(1, "Theo chiều dọc"), (1, "Theo chiều dọc"), (3, "Ngang (chi nhánh)"),
                                (3, "Ngang (sharding)")])]
    for i, (tieu_chi, o) in enumerate(hang):
        y = yb + i * rh
        if i % 2:
            hinh(s, MSO_SHAPE.RECTANGLE, MX + 0.05, y, c0 - 0.05, rh, "F3F5F9")
            for j in range(1, 4):
                hinh(s, MSO_SHAPE.RECTANGLE, MX + c0 + j * cw + 0.04, y, cw - 0.08, rh, "F3F5F9")
        chu(s, MX + 0.25, y, c0 - 0.3, rh, tieu_chi, 13, INK, True, neo="m")
        for j, (diem, mo_ta) in enumerate(o):
            x, c = MX + c0 + j * cw, mo_hinh[j][2]
            for k in range(3):
                hinh(s, MSO_SHAPE.OVAL, x + 0.18 + k * 0.17, y + rh / 2 - 0.06, 0.12, 0.12, c if k < diem else "D5DCE8")
            chu(s, x + 0.75, y, cw - 0.8, rh, mo_ta, 12, INK, neo="m")
    y = yb + 5 * rh
    chu(s, MX + 0.25, y, c0 - 0.3, dh, "Số liệu đo", 13, INK, True, neo="m")
    for j, so_lieu in enumerate(["Nguồn gốc, cài đặt đầy đủ", "Nạp lười: 889 câu SQL, nạp trước: 1",
                                 "Nửa kết giảm 73% byte truyền", "Hồ sơ 0,70 ms so với 3,62 ms"]):
        chu(s, MX + c0 + j * cw + 0.18, y, cw - 0.3, dh, so_lieu, 12, DAM[mo_hinh[j][2]], True, neo="m")
    y += dh + 0.15
    sh = hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, y, CW, BOT - y, NAVY, bo=0.1)
    chu_hinh(sh, "**Kết luận:** Quan hệ tập trung làm lõi · phân tán khi mở rộng ra nhiều tỉnh · NoSQL làm kho phụ "
             "(polyglot persistence)", 14, WHITE, False, nhan="FFB4AE")


# ----------------------------------------------------------------------------- 13. kiểm thử, CI
def kiem_thu():
    s = noi_dung(6, "Kiểm thử tích hợp: Dựng lại toàn bộ từ CSDL trống", "Như", 30, [
        "Toàn bộ chạy lại từ CSDL trống trong 1–2 phút, tự động bằng GitHub Actions khi có Pull Request hoặc push vào main.",
        "114 ca kiểm thử SQL (39 procedure, 22 trigger, 53 phân quyền) và 125 ca kiểm thử giao diện đều đạt."])
    ag = KHE_MUI_TEN
    bw = (CW - 4 * ag) / 5
    for i, (ic, ten, kq) in enumerate([("container", "Khởi động\nSQL Server 2025", "≈ 16 giây"),
                                       ("database", "Dựng CSDL\ntừ 9 script", "≈ 18–20 giây"),
                                       ("flask-conical", "Kiểm thử SQL: 4 bộ", "114 ca đạt"),
                                       ("search", "31 truy vấn\nnghiệp vụ", "≈ 4 giây, không lỗi"),
                                       ("layout-dashboard", "Kiểm thử giao diện", "125 ca đạt, ≈ 30 giây")]):
        x = MX + i * (bw + ag)
        the(s, x, TOP, bw, 1.95)
        bieu_tuong(s, ic, x + 0.2, TOP + 0.2, 0.5, NAVY)
        chu(s, x + 0.82, TOP + 0.2, bw - 0.9, 0.5, f"BƯỚC {i + 1}", 10, MUTED, True, neo="m", gian_chu=1)
        chu(s, x + 0.2, TOP + 0.85, bw - 0.4, 0.55, ten, 14, INK, True)
        chu(s, x + 0.2, TOP + 1.45, bw - 0.4, 0.3, kq, 13, DAM[GREEN], True)
        if i < 4:
            mui_ten(s, x + bw + (ag - 0.3) / 2, TOP + 0.82, 0.3)
    y, lw = TOP + 2.25, 7.4
    tw = (lw - 0.5) / 3
    for i, (so, nhan, mo_ta, c) in enumerate([("114", "ca kiểm thử SQL", "39 procedure · 22 trigger · 53 phân quyền", GREEN),
                                              ("125", "ca kiểm thử giao diện", "Mọi trang × mọi vai trò, xuất PDF, Excel, Word",
                                               GREEN),
                                              ("< 150 ms", "mỗi thao tác tra cứu", "Đo trên bản sao 10.121 học viên", NAVY)]):
        x = MX + i * (tw + 0.25)
        the(s, x, y, tw, BOT - y)
        y0 = y + (BOT - y - 1.95) / 2
        chu(s, x + 0.25, y0, tw - 0.5, 0.75, so, 34 if len(so) < 5 else 28, DAM[c], True, neo="m")
        chu(s, x + 0.25, y0 + 0.8, tw - 0.5, 0.3, nhan, 14, INK, True)
        chu(s, x + 0.25, y0 + 1.15, tw - 0.5, 0.8, mo_ta, 13, MUTED)
    x, w = MX + lw + 0.3, CW - lw - 0.3
    the(s, x, y, w, BOT - y)
    chu(s, x + 0.3, y + 0.22, w - 0.6, 0.32, "Quy trình làm việc nhóm", 15, INK, True)
    yy = y + 0.75
    for ic, mo_ta in [("git-pull-request", "30 công việc = 30 issue, mỗi việc một Pull Request vào develop"),
                      ("workflow", "GitHub Actions dựng lại và kiểm thử khi có PR hoặc push vào main"),
                      ("route", "Rà soát theo bài giảng → đợt khắc phục T30, đo lại số liệu")]:
        bieu_tuong(s, ic, x + 0.3, yy, 0.42, BLUE)
        chu(s, x + 0.9, yy - 0.05, w - 1.15, 0.55, mo_ta, 13, INK, neo="m")
        yy += 0.72


# ----------------------------------------------------------------------------- 14. demo
def demo():
    s = slide_moi(NAVY)
    dau_trang(s, 6, "DEMO TRỰC TIẾP", "Một vòng nghiệp vụ qua năm vai trò – khoảng 3 phút", toi=True)
    chan_trang(s, "Cả nhóm", toi=True)
    ghi_chu(s, "Cả nhóm", 180, [
        "Chuẩn bị sẵn 5 tab trình duyệt đã đăng nhập 5 vai trò, CSDL vừa dựng lại; mỗi bước khoảng 35 giây.",
        "Giáo vụ: Ghi danh vào lớp đầy để thấy lỗi từ CSDL. Kế toán: Lập phiếu thu 2 khoản rồi in PDF A5.",
        "Giáo viên: Điểm danh cả lớp, quét cảnh báo chuyên cần. Học viên: Chỉ thấy dữ liệu của mình.",
        "Quản lý: Dashboard, rồi mở nhật ký hệ thống để thấy các thao tác vừa làm. Thiếu thời gian thì bỏ bước 4."], 6)
    g = KHE_MUI_TEN
    cw, y, h = (CW - 4 * g) / 5, TOP + 0.2, 3.9
    for i, (vai, ic, viec, diem) in enumerate([
            ("Giáo vụ", "clipboard-pen", "Ghi danh học viên vào lớp đã đầy", "CSDL chặn: “Lớp đã đủ sĩ số.”"),
            ("Kế toán", "receipt", "Lập phiếu thu hai khoản, in PDF A5", "Master–detail trong một giao dịch"),
            ("Giáo viên", "calendar-check", "Điểm danh cả lớp, quét cảnh báo chuyên cần",
             "Trigger tính kết quả, cursor cảnh báo"),
            ("Học viên", "graduation-cap", "Xem lịch học, điểm, học phí", "Chỉ thấy dữ liệu của chính mình"),
            ("Quản lý", "gauge", "Mở dashboard, rồi nhật ký hệ thống", "Ai đã thao tác gì, lúc nào")]):
        x = MX + i * (cw + g)
        the(s, x, y, cw, h, NAVY2, None, False)
        chu_hinh(hinh(s, MSO_SHAPE.OVAL, x + 0.22, y + 0.25, 0.5, 0.5, CORAL), str(i + 1), 18)
        bieu_tuong(s, ic, x + cw - 0.72, y + 0.25, 0.5, BLUE)
        chu(s, x + 0.22, y + 0.95, cw - 0.44, 0.35, vai, 18, WHITE, True)
        chu(s, x + 0.22, y + 1.4, cw - 0.44, 1.2, viec, 14, WHITE)
        sh = the(s, x + 0.12, y + h - 1.15, cw - 0.24, 1.0, "2A4C86", None, False, 0.08)
        chu_hinh(sh, diem, 13, "F6C177", False, can="l", le=0.13)
        if i < 4:
            mui_ten(s, x + cw + (g - 0.3) / 2, y + h / 2 - 0.15, 0.3, BLUE)
    chu(s, MX, BOT - 0.3, CW, 0.3, "Mã nguồn: github.com/nhutruong-uit/ecms – dựng lại toàn bộ CSDL bằng một lệnh",
        13, PALE, neo="m")


# ----------------------------------------------------------------------------- 15. phân công
def phan_cong():
    s = noi_dung(6, "Phân công: 30 công việc, mỗi thành viên 5–7 việc", "Như", 20, [
        "Mỗi thành viên phụ trách 5–7 công việc trải đủ thiết kế, xử lý, giao diện và CSDL tiên tiến.",
        "Số commit chênh lệch vì nhóm trưởng làm thêm hạ tầng chung và đợt khắc phục T30; nhóm thống nhất mỗi người 20%."])
    cw = (CW - 1.0) / 5
    h = 4.25
    for i, (ten, chu_cai, c, vai, mssv, viec, mang) in enumerate([
            ("Trương Quang Như", "N", NAVY, "Nhóm trưởng", "25540022", "T01–T03, T15, T19, T29, T30",
             ["Quản lý dự án, Chương 1", "Hạ tầng, phân quyền", "Kiểm thử tự động, CI"]),
            ("Đỗ Phạm Minh Trâm", "T", BLUE, "Thành viên", "25540042", "T04–T07, T20, T25",
             ["ERD, sơ đồ lớp", "Chuẩn hoá, tạo bảng", "CSDL hướng đối tượng"]),
            ("Đỗ Bình Dương", "D", GREEN, "Thành viên", "25540008", "T08–T10, T18, T21, T27",
             ["Dữ liệu mẫu, truy vấn", "XML, nhập / xuất", "NoSQL (MongoDB)"]),
            ("Nguyễn Việt Phú", "P", CORAL, "Thành viên", "25540025", "T11–T14, T22, T26",
             ["Procedure, function", "Trigger, cursor", "CSDL phân tán"]),
            ("Nguyễn Bảo Giang", "G", VIOLET, "Thành viên", "25540009", "T16, T17, T23, T24, T28",
             ["View, sao lưu, phục hồi", "Báo cáo, dashboard", "Trang Tài chính"])]):
        x = MX + i * (cw + 0.25)
        the(s, x, TOP, cw, h)
        chu_hinh(hinh(s, MSO_SHAPE.OVAL, x + (cw - 0.85) / 2, TOP + 0.25, 0.85, 0.85, c), chu_cai, 28)
        chu(s, x + 0.1, TOP + 1.2, cw - 0.2, 0.32, ten, 15, INK, True, can="c")
        chu(s, x + 0.1, TOP + 1.5, cw - 0.2, 0.28, f"{vai} · {mssv}", 12, MUTED, can="c")
        nhan_tron(s, x + 0.1, TOP + 1.9, cw - 0.2, 0.3, viec, NHAT[c], DAM[c], 10.5)
        chu(s, x + 0.15, TOP + 2.4, cw - 0.3, 1.1, mang, 13, INK, can="c", sau=5)
        chu(s, x, TOP + h - 0.62, cw, 0.45, "20%", 24, CORAL, True, can="c", neo="m")
    chu(s, MX, BOT - 0.35, CW, 0.35, "141 commit  ·  30 issue, 30 Pull Request  ·  kế hoạch công khai tại "
        "nhutruong-uit.github.io/ecms", 13, MUTED, can="c", neo="m")


# ----------------------------------------------------------------------------- 16. kết luận
def ket_luan():
    s = noi_dung(6, "Kết luận và hướng phát triển", "Như", 40, [
        "Thông điệp chính: Luật nghiệp vụ nằm trong CSDL nên không ứng dụng nào bỏ qua được.",
        "Nêu thẳng hạn chế: Chưa kiểm thử tải, phân tán phần lớn mô phỏng, các tính năng bảo mật nâng cao mới chạy thử.",
        "Hướng phát triển nâng hệ thống lên tầng Trí tuệ của DIKW: Dự báo học viên bỏ học."])
    sh = hinh(s, MSO_SHAPE.ROUNDED_RECTANGLE, MX, TOP, CW, 0.85, NAVY, bo=0.12)
    bieu_tuong(s, "lightbulb", MX + 0.25, TOP + 0.175, 0.5, CORAL)
    chu(s, MX + 0.95, TOP, CW - 1.2, 0.85, "Luật nghiệp vụ nằm trong CSDL nên **không thể bị bỏ qua** – dù từ giao diện "
        "hay từ bất kỳ ứng dụng nào kết nối vào", 16, WHITE, neo="m", nhan="FFB4AE")
    del sh
    y, cw = TOP + 1.1, (CW - 0.6) / 3
    for i, (ic, c, ten, muc) in enumerate([
            ("circle-check", GREEN, "Đạt được", ["Đáp ứng đủ 22 yêu cầu chức năng",
                                                 "Đủ 5 giai đoạn: Thu thập → tổ chức → chuyển đổi → trình bày → an toàn",
                                                 "114 ca kiểm thử SQL, 125 ca kiểm thử giao diện đều đạt",
                                                 "Dựng lại từ CSDL trống bằng một lệnh, có CI"]),
            ("triangle-alert", AMBER, "Hạn chế", ["Chưa kiểm thử tải nhiều người dùng đồng thời",
                                                  "Streamlit tải lại cả trang, chưa tối ưu điện thoại",
                                                  "Phân tán phần lớn mô phỏng trên một máy",
                                                  "TDE, Audit, ledger mới chạy thử trên CSDL demo"]),
            ("rocket", BLUE, "Hướng phát triển", ["Cổng học viên, phụ huynh; thanh toán trực tuyến",
                                                  "REST API, ứng dụng di động, giữ lớp procedure",
                                                  "Redis bộ nhớ đệm, MongoDB đồng bộ bằng CDC",
                                                  "Kho dữ liệu, dự báo học viên bỏ học"])]):
        x = MX + i * (cw + 0.3)
        the(s, x, y, cw, BOT - y)
        tieu_de_the(s, x + 0.3, y + 0.25, cw - 0.6, ic, c, ten, d=0.5)
        yy = y + 1.0
        for m in muc:
            hinh(s, MSO_SHAPE.OVAL, x + 0.32, yy + 0.08, 0.1, 0.1, c)
            hh = cao(m, 13.5, cw - 0.85)
            chu(s, x + 0.55, yy, cw - 0.85, hh + 0.05, m, 13.5, INK)
            yy += hh + 0.18


# ----------------------------------------------------------------------------- 17. cảm ơn
def cam_on():
    s = slide_moi(NAVY)
    vong_trang_tri(s)
    chu(s, MX, 2.0, CW, 0.9, "Cảm ơn thầy và các bạn đã lắng nghe", 40, WHITE, True, can="c", neo="m")
    chu(s, MX, 2.95, CW, 0.55, "HỎI & ĐÁP", 18, CORAL, True, can="c", neo="m", gian_chu=3)
    for i, (ic, lien_ket) in enumerate([("git-branch", "github.com/nhutruong-uit/ecms"),
                                        ("route", "nhutruong-uit.github.io/ecms")]):
        w = 0.45 + 0.15 + rong(lien_ket, 16) + 0.35
        x = W / 2 - 0.15 - w if i == 0 else W / 2 + 0.15
        the(s, x, 4.15, w, 0.62, NAVY2, None, False, 0.31)
        bieu_tuong(s, ic, x + 0.1, 4.235, 0.45, BLUE)
        chu(s, x + 0.7, 4.15, w - 0.75, 0.62, lien_ket, 16, WHITE, neo="m")
    chu(s, MX, 6.55, CW, 0.3, "Nhóm 1 · IE103.Q21.VB2 · Giảng viên hướng dẫn: TS. Võ Phương Bình", 12, PALE2, can="c")
    ghi_chu(s, "Như (cả nhóm trả lời)", 10, [
        "Câu hỏi đã chuẩn bị: Vì sao đặt luật nghiệp vụ trong CSDL; vì sao chọn SQL Server thay vì MongoDB; "
        "SESSION_CONTEXT có an toàn khi dùng chung một login; khi nào nên dùng cursor; RPO, RTO đạt bao nhiêu.",
        "Người trả lời: Phần nào người phụ trách phần đó trả lời."])


def ve_bieu_tuong(src: Path):
    ICON.mkdir(parents=True, exist_ok=True)
    for ten in BIEU_TUONG:
        svg = (src / f"{ten}.svg").read_text().replace('stroke="currentColor"', 'stroke="#FFFFFF"')
        svg = svg.replace('stroke-width="2"', 'stroke-width="2.1"')
        subprocess.run(["rsvg-convert", "-w", "160", "-h", "160", "-o", str(ICON / f"{ten}.png")], input=svg.encode(),
                       check=True)
    print(f"✓ {len(BIEU_TUONG)} biểu tượng → {ICON.relative_to(ROOT)}")


def main():
    for dung in (bia, muc_luc, bai_toan, yeu_cau, erd, luoc_do, truy_van, an_ninh, lap_trinh, giao_dien, bao_cao,
                 tien_tien, kiem_thu, demo, phan_cong, ket_luan, cam_on):
        dung()
    assert STT["n"] == TONG, f"{STT['n']} slide, TONG = {TONG}"
    assert STT["phan"] == [p[4] for p in PHAN], f"thời lượng từng phần lệch: {STT['phan']}"
    assert STT["icon"] <= set(BIEU_TUONG), f"thiếu trong BIEU_TUONG: {STT['icon'] - set(BIEU_TUONG)}"
    for n, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            assert (sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= prs.slide_width
                    and sh.top + sh.height <= prs.slide_height), f"slide {n}: {sh.name} nằm ngoài slide"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"✓ {OUT.relative_to(ROOT)} ({STT['n']} slide, {STT['giay'] // 60}:{STT['giay'] % 60:02d} phút)")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--icons":
        ve_bieu_tuong(Path(sys.argv[2]))
    else:
        main()
