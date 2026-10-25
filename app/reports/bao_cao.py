"""Sinh báo cáo PDF / Excel / Word (T24) – thay cho Crystal Report của Bài thực hành 4.

Cấu trúc một báo cáo: Report Header (tên trung tâm, tiêu đề, tham số) → Page Header (tiêu đề cột lặp lại mỗi trang)
→ Detail (dòng dữ liệu) / Group Header – Group Footer (tổng phụ, thống kê nhóm) → Report Footer (tổng cộng,
trung bình / thấp nhất / cao nhất) → Page Footer ("Trang X / Y", thời điểm in).
Định dạng có điều kiện: hàm to_mau(dòng) → "do" | "xanh" | None tô màu dòng trong cả PDF, Excel và Word.
"""
from datetime import datetime
from io import BytesIO
from typing import Callable
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from reports.pdf_common import FONT, FONT_BOLD, money, styles

XANH = colors.HexColor("#1d4ed8")
NHAT = colors.HexColor("#f1f5f9")
NHOM = colors.HexColor("#dbeafe")
MAU = {"do": (colors.HexColor("#fee2e2"), colors.HexColor("#991b1b"), "FEE2E2", "991B1B"),
       "xanh": (colors.HexColor("#dcfce7"), colors.HexColor("#14532d"), "DCFCE7", "14532D")}
HAM = {"sum": ("Cộng", "sum", "SUM"), "avg": ("Trung bình", "mean", "AVERAGE"), "min": ("Thấp nhất", "min", "MIN"),
       "max": ("Cao nhất", "max", "MAX"), "count": ("Số dòng", "count", "COUNT")}
Cot = tuple[str, str, float, str | None]          # (tên cột df, tiêu đề, độ rộng mm, kiểu: money | date | num | None)


class _CanvasDanhSo(canvas.Canvas):
    """Page Footer: 'Trang X / Y' – cần biết tổng số trang nên vẽ chân trang sau khi dựng xong toàn bộ."""

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._trang = []

    def showPage(self):
        self._trang.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        tong = len(self._trang)
        for t in self._trang:
            self.__dict__.update(t)
            w, _ = self._pagesize
            self.setFont(FONT, 7.5)
            self.setFillColor(colors.grey)
            self.drawString(12 * mm, 8 * mm, f"ECMS · in lúc {datetime.now():%d/%m/%Y %H:%M}")
            self.drawRightString(w - 12 * mm, 8 * mm, f"Trang {self._pageNumber} / {tong}")
            super().showPage()
        super().save()


TIEN = ("HocPhi", "ThanhTien", "DaThu", "ConNo", "DoanhThu", "SoTien", "TongLuong", "LuongCoBan", "DonGiaBuoi")


def _fmt(v, kieu: str | None):
    if v is None or v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
        return ""
    if kieu == "money":
        return money(v)
    if kieu == "date":
        return pd.to_datetime(v).strftime("%d/%m/%Y")
    if kieu == "num":
        return f"{float(v):,.2f}".rstrip("0").rstrip(".").replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


def _thong_ke_rows(nhan_dau: str, d: pd.DataFrame, cot: list[Cot], thong_ke: dict[str, tuple[str, ...]]) -> list[list]:
    """Dòng thống kê (Trung bình / Thấp nhất / Cao nhất…) cho các cột trong thong_ke."""
    rows = []
    for h in dict.fromkeys(h for hs in thong_ke.values() for h in hs):
        r = [""] * len(cot)
        r[0] = f"{HAM[h][0]} {nhan_dau}".strip()
        for i, (c, _, _, k) in enumerate(cot):
            if h in thong_ke.get(c, ()):
                v = pd.to_numeric(d[c], errors="coerce")
                r[i] = _fmt(getattr(v, HAM[h][1])(), k)
        rows.append(r)
    return rows


def pdf(tieu_de: str, tham_so: str, df: pd.DataFrame, cot: list[Cot], nhom: str | None = None, tong: list[str] = (),
        ngang: bool = True, thong_ke: dict[str, tuple[str, ...]] | None = None,
        to_mau: Callable[[dict], str | None] | None = None) -> bytes:
    """cot: [(tên cột df, tiêu đề, độ rộng mm, kiểu)]. nhom: cột phân nhóm (Group Header/Footer). tong: cột cộng SUM.
    thong_ke: {cột: ("avg", "min", "max")} – thêm dòng thống kê ở chân nhóm và chân báo cáo. to_mau: định dạng có điều kiện."""
    st = styles()
    o = ParagraphStyle("o", fontName=FONT, fontSize=8.5, leading=10.5)                      # ô chữ: tự xuống dòng trong cột
    # TEXTCOLOR của TableStyle không áp vào Paragraph → ô chữ của dòng tô màu cần style riêng mang màu chữ
    o_mau = {m: ParagraphStyle(f"o_{m}", parent=o, textColor=v[1]) for m, v in MAU.items()}
    buf = BytesIO()
    size = landscape(A4) if ngang else A4
    doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title=tieu_de, author="ECMS")
    el = [Paragraph("TRUNG TÂM TIẾNG ANH ECMS", st["n"]), Paragraph(escape(tieu_de.upper()), st["title"]),
          Paragraph(escape(tham_so), st["center"]), Spacer(0, 4 * mm)]                         # Report Header
    head = [c[1] for c in cot]
    rows, kieu_dong = [head], []
    thong_ke = thong_ke or {}

    def o_du_lieu(v, k, mau=None):
        txt = _fmt(v, k)
        return txt if k in ("money", "num", "date") else Paragraph(escape(txt), o_mau.get(mau, o))

    def dong_tong(nhan, d):
        r = [""] * len(cot)
        r[0] = nhan
        for i, (c, _, _, k) in enumerate(cot):
            if c in tong:
                r[i] = _fmt(d[c].sum(), k)
        return r

    def chi_tiet(d):
        for r in d.to_dict("records"):
            mau = to_mau(r) if to_mau else None
            rows.append([o_du_lieu(r.get(c), k, mau) for c, _, _, k in cot])
            kieu_dong.append(("d", mau))

    if nhom:
        for g, d in df.groupby(nhom, sort=False):
            nhan = dict((c[0], c[1]) for c in cot).get(nhom)
            rows.append([f"{nhan}: {g}" if nhan else str(g)] + [""] * (len(cot) - 1)); kieu_dong.append(("gh", None))
            chi_tiet(d)
            if tong:
                rows.append(dong_tong(f"Cộng {g}", d)); kieu_dong.append(("gf", None))           # Group Footer
            for r in _thong_ke_rows("", d, cot, thong_ke):
                rows.append(r); kieu_dong.append(("gf", None))
    else:
        chi_tiet(df)
    if tong:
        rows.append(dong_tong("TỔNG CỘNG", df)); kieu_dong.append(("rf", None))                 # Report Footer
    for r in _thong_ke_rows("toàn báo cáo" if nhom else "", df, cot, thong_ke):
        rows.append(r); kieu_dong.append(("rf", None))
    kha_dung = size[0] - 24 * mm                                                                 # co cột cho vừa khổ giấy
    rong = [c[2] * mm for c in cot]
    he_so = min(1.0, kha_dung / sum(rong))
    t = Table(rows, colWidths=[w * he_so for w in rong], repeatRows=1)                           # Page Header lặp lại
    sty = [("FONT", (0, 0), (-1, -1), FONT, 8.5), ("FONT", (0, 0), (-1, 0), FONT_BOLD, 8.5),
           ("BACKGROUND", (0, 0), (-1, 0), XANH), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
           ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, (_, _, _, k) in enumerate(cot):
        if k in ("money", "num"):
            sty.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    for j, (kd, mau) in enumerate(kieu_dong, start=1):
        if kd == "gh":
            sty += [("SPAN", (0, j), (-1, j)), ("BACKGROUND", (0, j), (-1, j), NHOM), ("FONT", (0, j), (-1, j), FONT_BOLD, 9)]
        elif kd == "gf":
            sty += [("FONT", (0, j), (-1, j), FONT_BOLD, 8.5), ("LINEABOVE", (0, j), (-1, j), 0.8, XANH)]
        elif kd == "rf":
            sty += [("FONT", (0, j), (-1, j), FONT_BOLD, 9.5), ("BACKGROUND", (0, j), (-1, j), NHAT),
                    ("LINEABOVE", (0, j), (-1, j), 1.2, colors.black)]
        elif mau in MAU:                                                                        # định dạng có điều kiện
            sty += [("BACKGROUND", (0, j), (-1, j), MAU[mau][0]), ("TEXTCOLOR", (0, j), (-1, j), MAU[mau][1])]
        elif j % 2 == 0:
            sty.append(("BACKGROUND", (0, j), (-1, j), NHAT))
    t.setStyle(TableStyle(sty))
    el += [t, Spacer(0, 3 * mm), Paragraph(f"Số dòng: {len(df)}", st["small"])]
    doc.build(el, canvasmaker=_CanvasDanhSo)
    return buf.getvalue()


def excel(sheets: dict[str, tuple[str, pd.DataFrame, dict[str, str]]], tong: dict[str, list[str]] | None = None,
          tham_so: str | None = None, thong_ke: dict[str, dict[str, tuple[str, ...]]] | None = None,
          to_mau: dict[str, Callable[[dict], str | None]] | None = None) -> bytes:
    """sheets: {tên sheet: (tiêu đề, df, {cột: tiêu đề hiển thị})}. Định dạng tiền, ngày, tiêu đề, dòng tổng (công thức SUM),
    dòng thống kê (AVERAGE / MIN / MAX), tô màu có điều kiện, trang in A4 ngang lặp tiêu đề cột, chân trang "Trang X / Y"."""
    wb = Workbook()
    wb.remove(wb.active)
    vien = Border(*(Side(style="thin", color="CBD5E1"),) * 4)
    for ten, (tieu_de, df, nhan) in sheets.items():
        ws = wb.create_sheet(ten[:31])
        cols = list(nhan)
        ws.cell(1, 1, "TRUNG TÂM TIẾNG ANH ECMS").font = Font(bold=True, color="1D4ED8")
        ws.cell(2, 1, tieu_de).font = Font(bold=True, size=14)
        ws.cell(3, 1, (f"{tham_so} · " if tham_so else "") + f"In lúc {datetime.now():%d/%m/%Y %H:%M}").font = Font(
            italic=True, color="64748B")
        for j, c in enumerate(cols, 1):
            cell = ws.cell(5, j, nhan[c])
            cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D4ED8"), vien
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ham_mau = (to_mau or {}).get(ten)
        for i, rec in enumerate(df.to_dict("records"), 6):
            mau = ham_mau(rec) if ham_mau else None
            for j, c in enumerate(cols, 1):
                v = rec.get(c)
                v = None if (isinstance(v, float) and pd.isna(v)) or v is pd.NaT else v
                cell = ws.cell(i, j, float(v) if hasattr(v, "is_finite") else v)
                cell.border = vien
                if c in TIEN:
                    cell.number_format = '#,##0" đ"'
                if mau in MAU:                                                                  # định dạng có điều kiện
                    cell.fill, cell.font = PatternFill("solid", fgColor=MAU[mau][2]), Font(color=MAU[mau][3])
        cuoi = 5 + len(df)
        dong = cuoi
        if tong and ten in tong and len(df):
            dong += 1
            ws.cell(dong, 1, "TỔNG CỘNG").font = Font(bold=True)
            for c in tong[ten]:
                j = cols.index(c) + 1
                L = get_column_letter(j)
                cell = ws.cell(dong, j, f"=SUM({L}6:{L}{cuoi})")                                # công thức Excel, không ghi số cứng
                cell.font, cell.number_format = Font(bold=True), '#,##0" đ"'
        tk = (thong_ke or {}).get(ten, {})
        for h in dict.fromkeys(h for hs in tk.values() for h in hs):
            if not len(df):
                break
            dong += 1
            ws.cell(dong, 1, HAM[h][0]).font = Font(bold=True)
            for c, hs in tk.items():
                if h in hs:
                    L = get_column_letter(cols.index(c) + 1)
                    cell = ws.cell(dong, cols.index(c) + 1, f"={HAM[h][2]}({L}6:{L}{cuoi})")
                    cell.font, cell.number_format = Font(bold=True), '#,##0" đ"' if c in TIEN else '0.00'
        for j, c in enumerate(cols, 1):
            w = max([len(str(nhan[c]))] + [len(str(v)) for v in df[c].head(200)]) + 2
            ws.column_dimensions[get_column_letter(j)].width = min(45, max(9, w))
        ws.freeze_panes = "A6"
        ws.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{max(cuoi, 6)}"
        # trang in: A4 ngang, vừa 1 trang bề ngang, lặp dòng tiêu đề cột, chân trang số trang
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "5:5"
        ws.oddFooter.left.text, ws.oddFooter.right.text = "ECMS · &D &T", "Trang &P / &N"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def word(tieu_de: str, tham_so: str, df: pd.DataFrame, cot: list[Cot], tong: list[str] = (),
         to_mau: Callable[[dict], str | None] | None = None) -> bytes:
    """Báo cáo dạng Word (.docx) cho người cần chỉnh sửa / ký duyệt: cùng cấu trúc Report Header – bảng (dòng tiêu đề
    lặp lại mỗi trang) – Report Footer, chân trang "Trang X / Y" bằng trường PAGE / NUMPAGES của Word."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    def to_nen(cell, hex6):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex6)
        cell._tc.get_or_add_tcPr().append(shd)

    def truong(p, ma):
        for kieu, txt in (("begin", None), (None, ma), ("end", None)):
            r = p.add_run()
            if kieu:
                f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kieu); r._r.append(f)
            else:
                t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = txt; r._r.append(t)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation, sec.page_width, sec.page_height = WD_ORIENT.LANDSCAPE, Mm(297), Mm(210)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Mm(12))
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(9)
    doc.add_paragraph().add_run("TRUNG TÂM TIẾNG ANH ECMS").bold = True                         # Report Header
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(tieu_de.upper()); r.bold, r.font.size = True, Pt(15)
    p = doc.add_paragraph(tham_so); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = doc.add_table(rows=1, cols=len(cot)); t.style = "Table Grid"
    for i, (_, nhan, _, _) in enumerate(cot):
        c = t.rows[0].cells[i]; c.text = nhan; to_nen(c, "1D4ED8")
        c.paragraphs[0].runs[0].bold = True; c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true")                             # Page Header lặp lại mỗi trang
    t.rows[0]._tr.get_or_add_trPr().append(hdr)
    for rec in df.to_dict("records"):
        mau = to_mau(rec) if to_mau else None
        cells = t.add_row().cells
        for i, (c, _, _, k) in enumerate(cot):
            cells[i].text = _fmt(rec.get(c), k)
            if k in ("money", "num"):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if mau in MAU:
                to_nen(cells[i], MAU[mau][2])
    if tong:
        cells = t.add_row().cells
        cells[0].text = "TỔNG CỘNG"
        for i, (c, _, _, k) in enumerate(cot):
            if c in tong:
                cells[i].text = _fmt(df[c].sum(), k); cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for c in cells:
            for run in c.paragraphs[0].runs:
                run.bold = True
    doc.add_paragraph(f"Số dòng: {len(df)}")
    p = sec.footer.paragraphs[0]                                                                # Page Footer
    p.add_run(f"ECMS · in lúc {datetime.now():%d/%m/%Y %H:%M}        Trang ")
    truong(p, "PAGE"); p.add_run(" / "); truong(p, "NUMPAGES")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
