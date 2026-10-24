"""Kiểm thử đơn vị cho phần sinh báo cáo (không cần CSDL): số tiền bằng chữ, định dạng, PDF / Excel / Word hợp lệ.

    pytest tests/test_reports.py -q
"""
import sys
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "app"))

from reports import bao_cao  # noqa: E402
from reports.pdf_common import money, so_thanh_chu  # noqa: E402
from reports.phieu_thu_pdf import phieu_thu_pdf  # noqa: E402


@pytest.mark.parametrize("so,chu", [
    (0, "Không đồng"),
    (5_000_000, "Năm triệu đồng"),
    (1_250_000, "Một triệu hai trăm năm mươi nghìn đồng"),
    (21_005_000, "Hai mươi mốt triệu không trăm lẻ năm nghìn đồng"),
    (15_275_000, "Mười lăm triệu hai trăm bảy mươi lăm nghìn đồng"),
])
def test_so_thanh_chu(so, chu):
    assert so_thanh_chu(so) == chu


def test_dinh_dang():
    assert money(1234500) == "1.234.500 đ"
    assert bao_cao._fmt(7.5, "num") == "7,5"
    assert bao_cao._fmt(None, "money") == ""
    assert bao_cao._fmt(pd.Timestamp("2026-10-15"), "date") == "15/10/2026"


DF = pd.DataFrame({"MaHV": ["HV0001", "HV0002", "HV0003"], "HoTen": ["Nguyễn <Văn> A & B", "Trần Thị B", "Lê C"],
                   "DiemTB": [8.5, 4.0, None], "ThanhTien": [5_000_000, 3_500_000, 0], "KetQua": ["Đạt", "Không đạt", None]})
COT = [("MaHV", "Mã", 20, None), ("HoTen", "Họ tên", 60, None), ("DiemTB", "TB", 20, "num"), ("ThanhTien", "Thành tiền", 30, "money")]
MAU = lambda r: "do" if r.get("KetQua") == "Không đạt" else None  # noqa: E731


def test_pdf_co_thong_ke_va_to_mau():
    pdf = bao_cao.pdf("Bảng điểm <thử>", "Lớp A & B", DF, COT, tong=["ThanhTien"],
                      thong_ke={"DiemTB": ("avg", "min", "max")}, to_mau=MAU)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 2000            # ký tự < > & trong dữ liệu không làm hỏng PDF


def test_excel_co_cong_thuc_tham_so_va_trang_in():
    from openpyxl import load_workbook
    x = bao_cao.excel({"BD": ("Bảng điểm", DF, {c[0]: c[1] for c in COT})}, tong={"BD": ["ThanhTien"]}, tham_so="Lớp thử",
                      thong_ke={"BD": {"DiemTB": ("avg", "max")}}, to_mau={"BD": MAU})
    ws = load_workbook(BytesIO(x))["BD"]
    o = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=SUM(D6:D8)" in o and "=AVERAGE(C6:C8)" in o and "=MAX(C6:C8)" in o
    assert ws.cell(3, 1).value.startswith("Lớp thử")
    assert ws.page_setup.orientation == "landscape" and ws.print_title_rows.replace("$", "") == "5:5"
    assert ws.oddFooter.right.text == "Trang &P / &N"
    assert ws.cell(7, 1).fill.fgColor.rgb.endswith("FEE2E2")      # dòng "Không đạt" tô đỏ


def test_word_co_bang_va_so_trang():
    from docx import Document
    d = Document(BytesIO(bao_cao.word("Bảng điểm", "Lớp thử", DF, COT, tong=["ThanhTien"], to_mau=MAU)))
    t = d.tables[0]
    assert len(t.rows) == 1 + len(DF) + 1                           # tiêu đề + chi tiết + tổng cộng
    assert t.rows[-1].cells[3].text == "8.500.000 đ"
    assert "NUMPAGES" in d.sections[0].footer._element.xml


def test_phieu_thu_escape():
    pt = {"SoPhieu": "PT000001", "NgayThu": "2026-10-15 09:00", "TenCN": "ECMS Thủ Đức", "DiaChiCN": "KP6 & P. Linh Trung",
          "DienThoaiCN": "028", "HoTenHV": "<script>", "MaHV": "HV0001", "SoDienThoai": "0901", "DiaChiHV": None,
          "HinhThuc": "Tiền mặt", "GhiChu": "<KM10> giảm & tặng", "TongTien": 1_000_000, "NguoiLap": "Kế toán"}
    ct = pd.DataFrame([{"MaLop": "IE55-2601", "TenKH": "IELTS 5.5", "HocPhi": 1_000_000, "GiamGia": 0, "ThanhTien": 1_000_000,
                        "SoTien": 1_000_000, "ConNo": 0}])
    assert phieu_thu_pdf(pt, ct)[:5] == b"%PDF-"
